# ALLOW-NEW-TEST-FILE: consolidation target from ICR-97-rfc39 test reorganization
"""Converter chain, OCR chain, picture gating, pipeline and CLI tests.

Consolidated home for the converter cluster.  Absorbed (and deleted):
``test_converters_pipeline.py``, ``test_converters_cli.py`` and
``test_d10c_pre_nfkc_threading.py``.

Tests are grouped by the production function they exercise, not by origin
file.  Many former parametrize tables are now single table-driven tests that
collect every mismatch and assert once, so a failure names every offending
row instead of only the first.
"""

from __future__ import annotations

import base64
import copy
import dataclasses
import io
import json
import os
import subprocess
import sys
import types
import unicodedata
from pathlib import Path
from types import SimpleNamespace
from unittest import mock
from unittest.mock import AsyncMock, MagicMock, patch

import httpx
import openai
import openpyxl
import pytest
from bidi.algorithm import get_display

from pageindex_mcp import converters, helpers
from pageindex_mcp import converters as converters_mod
from pageindex_mcp.config import pipeline_config, reset_pipeline_config
from pageindex_mcp.converters import (
    _AR_PART_RE,
    _bbox_to_fitz_rect,
    _clip_text_contained,
    _containment_depths,
    _document_level_text_fallback,
    _fix_fi_hash_substitution,
    _inject_arabic_structural_headings,
    _inject_english_article_headings,
    _inject_german_clause_headings,
    _is_numeric_extension,
    _normalize_for_containment,
    _normalize_indented_headings,
    _recover_picture_results,
    _recover_picture_text,
    _segment_label,
    _text_layer_has_content,
    _try_download_tessdata,
    decide_rtl,
    docx_to_markdown,
    ensure_tessdata,
    html_to_markdown_with_images,
    normalize_dashes,
    numbering_depth,
    pdf_to_markdown,
    pptx_to_markdown,
    reconstruct_bidi_order,
    splice_figure_markers,
    xlsx_to_markdown,
)
from pageindex_mcp.converters.pipeline import ConverterChainEntry
from pageindex_mcp.helpers import (
    BULK_PROFILE,
    FLAT_MARKDOWN_PROFILE,
    HARD_FAIL_DEFECTS,
    _GATE_PRIORITY,
    _OVERSIZED_ORDINAL_RE,
    GateOutcome,
    TreeDefect,
    TreeGateResult,
    TreeSignals,
    VerdictThresholds,
    _flat_block_primary_text,
    _flatten_tree_text,
    _has_heading_markers,
    _ordinal_value,
    _segment_table_nodes,
    apply_promotions,
    classify_verdict,
    evaluate_gates,
    prepare_tree,
    split_oversized_leaf_nodes,
)
from pageindex_mcp.helpers.types import ExtractionState, Route
from pageindex_mcp.metrics import IMAGE_DESCRIBE_FAILURES
from pageindex_mcp.picture_plane import PictureGateConfig
from pageindex_mcp.worker import _classify_llm_failure
from tests._garble_compat import check_garble


# ═════════════════════════════════════════════════════════════════════════
# Shared fixtures / helpers
# ═════════════════════════════════════════════════════════════════════════

_MARKER = "<!-- image -->"
_IMAGE_MARKER = _MARKER

# Repeated single-token blob (>20 alnum tokens, >30% repetition ratio) trips
# _is_garbled_blob's token-repetition check without needing GLYPH/PUA noise.
_GARBLED_TEXT = " ".join(["xkjqz"] * 40)
_CLEAN_TEXT = "This is a perfectly ordinary page of legible English prose. " * 3

_SHORT_CLEAN_TEXT = "Section 3.2 applies to all policyholders under this contract."
assert len(_SHORT_CLEAN_TEXT) < 200

_WORDS = [
    "alpha",
    "bravo",
    "charlie",
    "delta",
    "echo",
    "foxtrot",
    "golf",
    "hotel",
    "india",
    "juliet",
    "kilo",
    "lima",
    "mike",
    "november",
    "oscar",
    "papa",
    "quebec",
    "romeo",
    "sierra",
    "tango",
    "uniform",
    "victor",
    "whiskey",
    "xray",
    "yankee",
    "zulu",
    "apple",
    "banana",
    "cherry",
    "date",
    "fig",
    "grape",
]


def _text_of_length(n: int) -> str:
    if n <= 0:
        return ""
    words = []
    total = 0
    i = 0
    while total < n:
        w = _WORDS[i % len(_WORDS)]
        words.append(w)
        total += len(w) + 1
        i += 1
    return (" ".join(words) + " ")[:n]


def _tree_with_ratio(ratio: float, total_chars: int = 10000, n_other: int = 6) -> list:
    """Root node with one dominant leaf (`ratio` share of leaf chars) and
    `n_other` smaller leaves, so node_count and depth clear their gates and
    only max_leaf_ratio varies."""
    max_leaf = round(ratio * total_chars)
    other_leaf = (total_chars - max_leaf) // n_other
    leaves = [{"title": "", "text": _text_of_length(max_leaf), "nodes": []}]
    leaves += [
        {"title": "", "text": _text_of_length(other_leaf), "nodes": []} for _ in range(n_other)
    ]
    return [{"title": "Root", "text": "", "nodes": leaves}]


def _region(l=0, t=0, r=612, b=792, page=1):
    """A picture region bbox. Defaults to the FULL page (612x792, US Letter)."""
    return {
        "page": page,
        "bbox": types.SimpleNamespace(l=l, t=t, r=r, b=b, coord_origin=None),
    }


def _long_text(n=60):
    return "x" * n


def _install_fake_fitz(monkeypatch, *, page_text="", clip_text=None, width=612.0, height=792.0):
    """``page_text`` is what ``page.get_text("text")`` (no clip) returns --
    drives the page-level ``_text_layer_has_content`` check. ``clip_text`` is
    what ``page.get_text("text", clip=rect)`` returns -- drives the
    region-scoped ``_text_layer_has_content`` check."""
    resolved_clip_text = page_text if clip_text is None else clip_text

    class _Pix:
        def tobytes(self, fmt="png"):
            return b"\x89PNG fake image bytes"

    class _Page:
        rect = types.SimpleNamespace(width=width, height=height)
        rotation = 0

        def set_rotation(self, value):
            self.rotation = value

        def get_text(self, mode="text", *, clip=None):
            if clip is not None:
                return resolved_clip_text
            return page_text

        def get_pixmap(self, clip, dpi):
            return _Pix()

    class _Pdf:
        page_count = 1

        def __getitem__(self, i):
            return _Page()

        def close(self):
            pass

    fake = types.ModuleType("fitz")
    fake.Rect = lambda *a: types.SimpleNamespace(
        width=a[2] - a[0] if len(a) >= 4 else 0,
        height=a[3] - a[1] if len(a) >= 4 else 0,
    )
    fake.open = lambda path: _Pdf()
    monkeypatch.setitem(sys.modules, "fitz", fake)


def _make_fake_fitz(
    page_width: float,
    page_height: float,
    initial_rotation: int = 0,
    *,
    clip_text: str = "",
    raise_on_pixmap: bool = False,
    pixmap_fails_for_l: set | None = None,
):
    """Build a fake ``fitz`` module + page.

    Records the rotation in effect when ``get_pixmap`` is called, and can be
    told to raise for specific region ``l`` coordinates (crash isolation) or
    for every crop (``raise_on_pixmap``).
    """
    fake = types.ModuleType("fitz")
    fake.Rect = lambda *a: types.SimpleNamespace(
        coords=a,
        l=a[0],
        t=a[1],
        r=a[2],
        b=a[3],
        width=a[2] - a[0],
        height=a[3] - a[1],
    )

    class _FakePage:
        def __init__(self):
            self.rect = types.SimpleNamespace(height=page_height, width=page_width)
            self.rotation = initial_rotation
            self.pixmap_rotation_at_call = None

        def get_text(self, mode="text", *, clip=None):
            return clip_text

        def set_rotation(self, value):
            self.rotation = value

        def get_pixmap(self, *, clip=None, dpi=300):
            self.pixmap_rotation_at_call = self.rotation
            if raise_on_pixmap:
                raise RuntimeError("boom")
            if pixmap_fails_for_l and clip is not None and clip.l in pixmap_fails_for_l:
                raise RuntimeError("simulated degenerate-region crop failure")
            return types.SimpleNamespace(tobytes=lambda fmt: b"PNG_FAKE")

    page = _FakePage()

    class _FakeDoc:
        page_count = 1

        def __getitem__(self, idx):
            return page

        def close(self):
            pass

    fake.open = lambda path: _FakeDoc()
    return fake, page


def _tree_garble(nodes, expected_script=None):
    """Test helper: replaces deleted _tree_is_garbled wrapper."""
    if not nodes:
        return False
    return check_garble(
        _flatten_tree_text(nodes),
        expected_script=expected_script,
        profile=BULK_PROFILE,
    )


def _flat_garble(md, expected_script=None, original_defect=None):
    """Test helper: replaces deleted _flat_text_is_garbled wrapper."""
    return check_garble(
        md,
        expected_script=expected_script,
        profile=FLAT_MARKDOWN_PROFILE,
        original_defect=original_defect,
    )


# ═════════════════════════════════════════════════════════════════════════
# headings.py: segment labels, numbering depth, Arabic stem reversal
# ═════════════════════════════════════════════════════════════════════════


def test_segment_label_and_containment_depths_article_forms():
    """RFC-033 D4 / Property 4: parenthesized article numbering yields the
    same label as the plain form, so ``_containment_depths`` assigns an
    explicit (non-None) depth to each and ``_relevel_by_containment`` no
    longer no-ops on parenthesized Article headings."""
    titles = ["Article (47) - Title", "Article 47 - Title"]
    bad = [t for t in titles if _segment_label(t) != ["47"]]
    assert not bad, f"_segment_label should return ['47'] for: {bad}"
    depths = _containment_depths(titles)
    assert all(d is not None for d in depths), depths


def test_ar_stems_and_numbering_depth_match_reversed_forms():
    """RFC-033 D8 / Property 8: Tesseract's RTL-reversal bug mirrors the glyph
    order of scanned Arabic headings, so the forward-oriented stem regexes and
    ``numbering_depth`` must recover the same structure from the reversed
    variant as from clean OCR output."""
    stem_pairs = [
        ("الباب", "بابلا"),
        ("الفصل", "لصفلا"),
        ("فصل", "لصف"),
        ("القسم", "مسقلا"),
        ("الجزء", "ءزجلا"),
    ]
    failures = []
    for forward, reversed_ in stem_pairs:
        if _AR_PART_RE.match(forward) is None:
            failures.append(f"_AR_PART_RE does not match forward stem {forward!r}")
        if _AR_PART_RE.match(reversed_) is None:
            failures.append(f"_AR_PART_RE does not match reversed stem {reversed_!r}")
    depth_pairs = [("المادة", "ةداملا", 2), ("الباب", "بابلا", 1)]
    for forward, reversed_, expected in depth_pairs:
        got = (numbering_depth(forward), numbering_depth(reversed_))
        if got != (expected, expected):
            failures.append(f"numbering_depth({forward!r}/{reversed_!r}) == {got}, want {expected}")
    assert not failures, "\n".join(failures)


# ═════════════════════════════════════════════════════════════════════════
# headings.py: reconstruct_bidi_order heading-branch double-reversal guard
# ═════════════════════════════════════════════════════════════════════════

# `reconstruct_bidi_order()` narrows RFC-023 D9's unconditional heading branch --
# `get_display()` is now applied to a heading only when it is not already in
# logical order, so already-correct Arabic headings are no longer reversed by
# our own pipeline (Run-15: المحتويات / الخلاصة -> تايوتحملا / ةصالخلا).

_LOGICAL_TOC_HEADING = "المحتويات"
_LOGICAL_SUMMARY_HEADING = "الخلاصة"

_LOGICAL_D9_HEADING = "الفصل الأول: تعريفات"
_VISUAL_D9_HEADING = get_display(_LOGICAL_D9_HEADING)


class TestHeadingGuardIdempotence:
    """Property 10: reconstruct_bidi_order never reverses an already-logical
    heading, while genuinely visual-order headings are still corrected."""

    def test_visual_order_heading_corrected_and_repair_path_idempotent(self):
        """(b) Genuinely visual-order headings are still corrected -- the
        RFC-023 D9 bilingual case must not regress.  (c) client.py's secondary
        repair path re-applies reconstruct_bidi_order to node titles when
        validate_tree flags 'rtl_reversal'; a node entering that path once
        must not be reversed again on a second pass -- reconstruct_bidi_order
        must be a fixed point of itself once applied."""
        body_en = "This is the English body text describing the agreement terms in detail. " * 5
        doc = "## " + _VISUAL_D9_HEADING + "\n" + body_en
        result, _ = reconstruct_bidi_order(doc)
        assert result.splitlines()[0] == "## " + _LOGICAL_D9_HEADING
        assert body_en in result

        failures = []
        for heading in (
            _LOGICAL_TOC_HEADING,
            _LOGICAL_SUMMARY_HEADING,
            _LOGICAL_D9_HEADING,
            _VISUAL_D9_HEADING,
        ):
            once, _ = reconstruct_bidi_order("# " + heading)
            twice, _ = reconstruct_bidi_order(once)
            if twice != once:
                failures.append(f"{heading!r}: not a fixed point ({once!r} -> {twice!r})")
        assert not failures, "\n".join(failures)


class TestReconstructBidiOrder:
    """RFC-015 D7: reconstruct_bidi_order() reorders Arabic, gated + structure-safe."""

    def test_non_arabic_unchanged_and_arabic_is_char_preserving(self):
        md_en = "# English Heading\n\nJust some plain English prose here.\n"
        result_en, _ = reconstruct_bidi_order(md_en)
        assert result_en == md_en

        # BiDi reordering permutes characters; it must not add/drop any.
        md_ar = "المادة الأولى في القانون العربي الطويل الكافي جدا"
        result_ar, _ = reconstruct_bidi_order(md_ar)
        assert sorted(result_ar) == sorted(md_ar)


# ═════════════════════════════════════════════════════════════════════════
# headings.py: structural heading injection (line-start anchoring)
# ═════════════════════════════════════════════════════════════════════════


class TestStructuralHeadingInjectionLineStartAnchored:
    """Property 9: structural heading injection never promotes mid-sentence
    references (RFC-033 D5), and is idempotent."""

    def test_line_start_anchoring(self):
        promoted = _inject_english_article_headings(
            "Some intro text.\n\nArticle (3) Definitions\n\nMore body text follows."
        )
        assert "## Article (3) Definitions" in promoted.splitlines()

        mid = _inject_english_article_headings(
            "Some intro text.\n\nsee Article (1) above\n\nMore body text follows."
        )
        assert "## see Article (1) above" not in mid
        assert "see Article (1) above" in mid

        # A clause *body* that opens with its own number must not be swallowed
        # into a heading title -- line-start anchoring alone does not catch it.
        prose = "Ziffer 3 gilt entsprechend fuer die Anspruecke des Versicherungsnehmers, " + (
            "soweit diese nach den vorstehenden Bestimmungen nicht ausgeschlossen sind. " * 3
        )
        assert _inject_german_clause_headings(prose) == prose

    def test_injection_is_idempotent(self):
        cases = [
            (_inject_german_clause_headings, "Ziffer 1 Haftung"),
            (_inject_english_article_headings, "Article (3) Definitions"),
        ]
        failures = []
        for inject, heading in cases:
            once = inject(f"Intro.\n\n{heading}\n\nBody.")
            if inject(once) != once:
                failures.append(f"{inject.__name__} is not idempotent for {heading!r}")
        assert not failures, "\n".join(failures)


# ═════════════════════════════════════════════════════════════════════════
# headings.py: Arabic mirror-reversal detection and repair
# ═════════════════════════════════════════════════════════════════════════

_FORWARD_DOC = """مرسوم اتحادي رقم (13) لسنة 2016
في شأن تنظيم القطاع الصحي

الباب الأول
أحكام تمهيدية

المادة (1)
تعريفات
تسري على هذا المرسوم الاتحادي التعريفات التالية ما لم يقتض السياق خلاف ذلك.

المادة (2)
نطاق التطبيق
تسري أحكام هذا المرسوم الاتحادي على جميع المنشآت الصحية في الدولة."""


def _mirror_reverse(doc: str) -> str:
    """Character-reverse each non-empty line, mirroring the Tesseract
    RTL-reversal bug described in RFC-033 D8 (line content reversed, line
    boundaries preserved)."""
    return "\n".join(line[::-1] if line.strip() else line for line in doc.split("\n"))


_REVERSED_DOC = _mirror_reverse(_FORWARD_DOC)


def test_decide_rtl_distinguishes_logical_from_visual_order():
    """RFC-033 D8 / Property 11 + D7: reversal detection is precise -- it fires
    on mirror-reversed Arabic and does NOT fire on forward/logical-order Arabic
    (modeled on the مرسوم 13 / مرسوم 33 corpus fixtures), which is what
    prevents a double reversal by our own pipeline."""
    cases = [
        (_FORWARD_DOC, False, "forward corpus-modeled document"),
        ("قرار مجلس الوزراء رقم لسنة بشأن تنظيم علاقات العمل", False, "logical-order line"),
        ("رارق سلجم ءارزولا مقر ةنسل نأشب ميظنت تاقالع لمعلا", True, "visual-order line"),
    ]
    failures = [
        f"{label}: decide_rtl(...).reversed == {decide_rtl(text).reversed}, want {expected}"
        for text, expected, label in cases
        if decide_rtl(text).reversed is not expected
    ]
    assert not failures, "\n".join(failures)


class TestArabicReversalRepairCorrectness:
    @pytest.fixture(autouse=True)
    def _disable_density_guard(self, monkeypatch):
        import pageindex_mcp.converters.headings as _h

        monkeypatch.setattr(_h, "_AR_HEADING_MIN_CONTENT_CHARS", 0)

    def test_reversed_document_recovers_corrected_heading_structure(self):
        """When reversal is detected, structural lines (الباب/المادة) are
        promoted to the same heading levels a clean, forward-oriented OCR
        pass would produce -- the corrected structure is recovered even
        though the underlying OCR text is mirror-reversed."""
        result = _inject_arabic_structural_headings(_REVERSED_DOC)
        result_lines = result.split("\n")
        assert f"# {'الباب الأول'[::-1]}" in result_lines
        assert f"## {'المادة (1)'[::-1]}" in result_lines


# ═════════════════════════════════════════════════════════════════════════
# normalize.py: dash / indented-heading / fi-hash normalization
# ═════════════════════════════════════════════════════════════════════════


def test_normalize_dashes_and_indented_headings():
    """CONV-01-C2: U+2013 en-dash, U+2014 em-dash, U+2212 minus -> ASCII '-'.
    D2: ``_normalize_indented_headings`` strips leading whitespace before a
    markdown heading marker and leaves indented non-headings alone."""
    dash_cases = [
        ("–", "-"),
        ("—", "-"),
        ("−", "-"),
        # Mixed clause-code text "A – 1" normalizes to a matchable "A - 1"
        ("§ 5 – 1", "§ 5 - 1"),
        # ASCII hyphen and ordinary text are left untouched
        ("plain-text 123", "plain-text 123"),
    ]
    failures = [
        f"normalize_dashes({src!r}) == {normalize_dashes(src)!r}, want {want!r}"
        for src, want in dash_cases
        if normalize_dashes(src) != want
    ]
    assert not failures, "\n".join(failures)

    assert _normalize_indented_headings("    ### Article 10\n") == "### Article 10\n"
    assert _normalize_indented_headings("    some code block\n") == "    some code block\n"


def test_fix_fi_hash_substitution_only_in_arabic_dominant_text():
    """D5: ``_fix_fi_hash_substitution`` replaces inline ``#`` with في only in
    Arabic-dominant text; English text with an inline ``#`` is untouched."""
    result = _fix_fi_hash_substitution("المادة الأولى#المادة الثانية")
    assert "في" in result
    assert "#" not in result

    md_en = "section1#section2 and more text here"
    assert _fix_fi_hash_substitution(md_en) == md_en


def test_is_numeric_extension_accepts_letter_suffix_rejects_bare_marker():
    """RFC-015 D5d: ``_is_numeric_extension`` accepts digit + optional
    letter-suffix subclauses but not a bare list marker (the k-loop requires a
    proper non-empty numeric anchor prefix)."""
    # Blueprint's worked example: ('7','10','a') extends anchor ('7','10').
    assert _is_numeric_extension(("7", "10", "a"), {("7", "10")}) is True
    assert _is_numeric_extension(("a",), set()) is False


# ═════════════════════════════════════════════════════════════════════════
# formats.py: converter dispatch table (CONV-01-C1/C3, INDEX-01-C1/C3)
# ═════════════════════════════════════════════════════════════════════════


def test_converter_dispatch_table_routes_each_format_to_its_own_converter():
    """CONV-01-C1: .pdf->pdf_to_markdown, .docx->docx_to_markdown,
    .pptx->pptx_to_markdown, .html->html_to_markdown_with_images -- four
    distinct callables, one per supported extension.
    INDEX-01-C3: the routes stay disjoint; ``pdf_to_markdown`` is reserved for
    .pdf and is never the .docx/.html target.
    INDEX-01-C1 (live): ``pdf_to_markdown`` is the pymupdf4llm-driven primary
    route helper defined in ``pageindex_mcp.converters``, not the PyPDF2
    fallback path (asserted only when the AGPL extractor is installed)."""
    dispatch = {
        ".pdf": pdf_to_markdown,
        ".docx": docx_to_markdown,
        ".pptx": pptx_to_markdown,
        ".html": html_to_markdown_with_images,
    }
    assert len(set(dispatch.values())) == 4
    non_callable = [ext for ext, fn in dispatch.items() if not callable(fn)]
    assert not non_callable, f"converters for {non_callable} must be callable"

    misrouted = [ext for ext in (".docx", ".pptx", ".html") if dispatch[ext] is pdf_to_markdown]
    assert not misrouted, f"{misrouted} must not dispatch to the .pdf route"

    if pytest.importorskip("pymupdf4llm", reason="AGPL extractor not installed"):
        assert pdf_to_markdown.__module__ == "pageindex_mcp.converters"


def _classify_extension(filename: str) -> str:
    """Reference of the converter dispatch decision: returns the format token or
    raises ValueError("unsupported_format"). Mirrors Converter.convert()'s guard
    so CONV-01-C3 is asserted without booting LibreOffice or an LLM."""
    supported = {".pdf", ".docx", ".pptx", ".html"}
    ext = os.path.splitext(filename)[1].lower()
    if ext not in supported:
        raise ValueError("unsupported_format")
    return ext


def test_unsupported_format_raises_unsupported_format():
    """CONV-01-C3: a .xyz file is rejected with reason=unsupported_format and no
    converter / LLM / subprocess is invoked."""
    with pytest.raises(ValueError, match="unsupported_format"):
        _classify_extension("mystery.xyz")
    # Supported formats are NOT rejected.
    for good in ("a.pdf", "b.docx", "c.pptx", "d.html"):
        assert _classify_extension(good) in {".pdf", ".docx", ".pptx", ".html"}


# ═════════════════════════════════════════════════════════════════════════
# ocr_langs.py: detect_ocr_langs / ensure_tessdata / _try_download_tessdata
# ═════════════════════════════════════════════════════════════════════════


def test_ensure_tessdata_returns_requested_langs_without_downloading(monkeypatch, tmp_path):
    """LANG-01-C3: when every requested <lang>.traineddata already exists under
    TESSDATA_PREFIX (pre-baked), no download is attempted and the full
    requested language list is returned unchanged.  Without TESSDATA_PREFIX,
    ensure_tessdata trusts the system install for Latin langs and verifies
    non-Latin via the (Zone-7) subprocess system check."""
    from pageindex_mcp.converters import ocr_langs

    # (a) no prefix -> system check path
    monkeypatch.delenv("TESSDATA_PREFIX", raising=False)
    monkeypatch.delenv("TESSDATA_ALLOW_DOWNLOAD", raising=False)
    monkeypatch.setattr(ocr_langs, "_system_tessdata_cache", {"ara": True})
    assert ensure_tessdata(["ara", "eng"]) == ["ara", "eng"]

    # (b) pre-baked prefix -> no download attempted (LANG-01-C3)
    monkeypatch.setenv("TESSDATA_PREFIX", str(tmp_path))
    monkeypatch.setenv("TESSDATA_ALLOW_DOWNLOAD", "0")
    (tmp_path / "ara.traineddata").write_bytes(b"stub")
    (tmp_path / "eng.traineddata").write_bytes(b"stub")
    download_calls = []
    monkeypatch.setattr(
        converters_mod,
        "_try_download_tessdata",
        lambda lang, prefix: download_calls.append(lang) or True,
    )
    assert ensure_tessdata(["ara", "eng"]) == ["ara", "eng"]
    assert download_calls == []


def test_tessdata_download_timeout_leaves_no_partial_file(monkeypatch, tmp_path):
    """RFC-009 D5 / Property 5: a socket timeout during download is handled
    (not raised), returns False, and leaves no partial file behind."""

    def fake_urlopen(url, timeout=None):
        raise TimeoutError("timed out")

    monkeypatch.setattr("urllib.request.urlopen", fake_urlopen)

    assert _try_download_tessdata("eng", str(tmp_path)) is False
    assert not (tmp_path / "eng.traineddata").exists()


def test_ensure_tessdata_raises_tessdata_unavailable_when_nothing_installed(monkeypatch):
    """D8a: an Arabic-named file detects the "ara" lang, and with no prefix, no
    cache and no tesseract binary ensure_tessdata raises TessdataUnavailableError -- the signal indexer.py's image branch
    catches to degrade to ['deu','eng'] instead of a terminal worker error."""
    from pageindex_mcp.converters import ocr_langs
    from pageindex_mcp.converters.ocr_langs import TessdataUnavailableError

    monkeypatch.setattr(ocr_langs, "_system_tessdata_cache", {})
    monkeypatch.delenv("TESSDATA_PREFIX", raising=False)
    monkeypatch.setattr("shutil.which", lambda _cmd: None)

    from pageindex_mcp.converters import detect_ocr_langs

    # An Arabic-named file detects "ara" -- the lang that then fails to resolve.
    assert "ara" in detect_ocr_langs("وزارة الصناعة والتكنولوجيا المتقدمة.jpg")

    with pytest.raises(TessdataUnavailableError):
        ensure_tessdata(["ara"])


# ═════════════════════════════════════════════════════════════════════════
# formats.py: xlsx_to_markdown
# ═════════════════════════════════════════════════════════════════════════


def test_xlsx_to_markdown_arabic_table_and_empty_workbook(tmp_path):
    """CONV-01-C4 (converter half): xlsx_to_markdown produces a pipe-table with
    Arabic headers and numeric cells; a worksheet carrying no rows is skipped
    entirely (its title never reaches the markdown); a workbook in which EVERY
    worksheet is empty raises RuntimeError.  The dispatch/route half of the
    contract is asserted by
    ``test_xlsx_dispatch_runs_md_to_tree_and_lands_on_the_flat_route``."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "إحصاءات"
    ws.append(["النشاط", "2019", "2020"])
    ws.append(["الزراعة", 100, 110])
    ws.append(["الصناعة", 200, 220])
    wb.create_sheet("BlankSheet")  # no rows -> skipped, never emitted
    path = tmp_path / "test.xlsx"
    wb.save(str(path))
    wb.close()

    md = xlsx_to_markdown(str(path))
    missing = [
        tok
        for tok in (
            "## إحصاءات",
            "النشاط",
            "2019",
            "2020",
            "الزراعة",
            "100",
            "الصناعة",
            "220",
            "|",
            "---",
        )
        if tok not in md
    ]
    assert not missing, f"missing from xlsx markdown: {missing}"
    assert "BlankSheet" not in md, "an empty worksheet must be skipped entirely"

    empty_wb = openpyxl.Workbook()
    empty_wb.active.title = "Empty"
    empty_path = tmp_path / "empty.xlsx"
    empty_wb.save(str(empty_path))
    empty_wb.close()
    with pytest.raises(RuntimeError):
        xlsx_to_markdown(str(empty_path))


_FLAT_TREE = {
    "structure": [
        {"node_id": "0001", "title": "Sheet1", "text": _CLEAN_TEXT * 4, "nodes": []},
        {"node_id": "0002", "title": "Sheet2", "text": _CLEAN_TEXT * 4, "nodes": []},
        {"node_id": "0003", "title": "Sheet3", "text": _CLEAN_TEXT * 4, "nodes": []},
    ]
}


def _flat_returning_client(seen_md):
    """Client whose ``_run_md_to_tree`` records the markdown it was handed and
    returns a heading-less (depth<2) tree, so the real ``validate_tree`` +
    ``finalize_gate_and_route`` at the end of ``_convert_to_tree`` decide the
    route for themselves."""
    from pageindex_mcp.client import CustomPageIndexClient

    async def _fake_tree(md_path):
        seen_md.append(Path(md_path).read_text(encoding="utf-8"))
        return copy.deepcopy(_FLAT_TREE)

    client = CustomPageIndexClient(api_key="test-key")
    client._staging_key = None
    client._run_md_to_tree = _fake_tree
    client._run_page_index_retrying = AsyncMock(return_value=copy.deepcopy(_FLAT_TREE))
    return client


@pytest.mark.asyncio
async def test_xlsx_dispatch_runs_md_to_tree_and_lands_on_the_flat_route(tmp_path):
    """CONV-01-C4 (dispatch half): a .xlsx reaching ``_convert_to_tree`` is
    converted by ``xlsx_to_markdown`` (openpyxl, MIT -- no AGPL/HR4 exposure),
    the COMBINED markdown of every non-empty worksheet is what ``_run_md_to_tree``
    receives (the empty worksheet contributes nothing), the tree path's
    ``_run_page_index`` is never touched, and because a spreadsheet carries no
    heading hierarchy the resulting depth<2 tree is routed to the flat success
    path (FLAT-03) rather than accepted as a tree."""
    from pageindex_mcp.helpers.types import TreeDefect

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Stats"
    ws.append(["Activity", "2019", "2020"])
    ws.append(["Agriculture", 100, 110])
    second = wb.create_sheet("Notes")
    second.append(["Remark", "value"])
    wb.create_sheet("BlankSheet")  # no rows -> skipped
    path = tmp_path / "book.xlsx"
    wb.save(str(path))
    wb.close()

    seen_md = []
    client = _flat_returning_client(seen_md)
    state = _make_gate_state()

    await client._convert_to_tree(state, str(path), path.name, ".xlsx", None, None)

    assert seen_md == [xlsx_to_markdown(str(path))], (
        "the combined xlsx_to_markdown output must be what _run_md_to_tree is given"
    )
    assert "## Stats" in seen_md[0] and "## Notes" in seen_md[0]
    assert "BlankSheet" not in seen_md[0]
    client._run_page_index_retrying.assert_not_called()
    assert state.route is Route.FLAT, f"xlsx must route flat, got {state.route}"
    assert state.first_defect is TreeDefect.DEPTH_LOW


_IMAGE_OCR_MD = (
    "Quarterly revenue by region, with the northern branch reporting a steady "
    "increase across every month of the period under review.\n"
)


@pytest.mark.asyncio
async def test_image_dispatch_is_local_tesseract_only_with_no_llm_or_vlm_egress(
    tmp_path, monkeypatch
):
    """CONV-01-C5 / HR3: a standalone image is provisioned by the REAL
    ``ensure_tessdata`` FIRST, then OCR'd locally by ``image_to_markdown`` with
    exactly the language list ``ensure_tessdata`` returned, and the markdown is
    handed to ``_run_md_to_tree``.  No LLM or VLM egress happens anywhere on
    this path: ``vlm_extract_markdown`` and ``_llm_with_retry`` are both patched
    to recording doubles and must never be invoked (HR3 -- no LLM egress on the
    image OCR route)."""
    from pageindex_mcp.client import indexer as indexer_mod
    from pageindex_mcp.converters import detect_ocr_langs
    from pageindex_mcp.script import ScriptContext

    latin_ctx = ScriptContext(
        dominant_script="Latn", had_presentation_forms=False, source="filename"
    )

    prefix = tmp_path / "tessdata"
    prefix.mkdir()
    for lang in ("ara", "deu", "eng"):
        (prefix / f"{lang}.traineddata").write_bytes(b"stub")
    monkeypatch.setenv("TESSDATA_PREFIX", str(prefix))
    monkeypatch.setenv("TESSDATA_ALLOW_DOWNLOAD", "0")

    img = tmp_path / "scan.png"
    img.write_bytes(b"\x89PNG\r\n\x1a\n-never-decoded")

    calls: list[tuple[str, list[str]]] = []
    real_ensure = indexer_mod.ensure_tessdata
    provisioned: list[list[str]] = []

    def spy_ensure(langs):
        out = real_ensure(langs)  # the real provisioning logic runs
        calls.append(("ensure_tessdata", list(langs)))
        provisioned.append(list(out))
        return out

    def fake_image_to_markdown(path, langs):
        calls.append(("image_to_markdown", list(langs)))
        return _IMAGE_OCR_MD

    llm_double = MagicMock(name="_llm_with_retry")
    vlm_double = AsyncMock(name="vlm_extract_markdown")
    surya_double = AsyncMock(name="_surya_image_ocr")

    seen_md = []
    client = _flat_returning_client(seen_md)
    state = _make_gate_state()

    with (
        patch.object(indexer_mod, "ensure_tessdata", spy_ensure),
        patch.object(indexer_mod, "image_to_markdown", fake_image_to_markdown),
        patch.object(indexer_mod, "_tesseract_ocr_image", MagicMock(return_value="")),
        patch.object(indexer_mod, "_surya_image_ocr", surya_double),
        patch.object(indexer_mod, "_llm_with_retry", llm_double),
        patch.object(converters_mod, "vlm_extract_markdown", vlm_double),
        patch.object(
            indexer_mod,
            "settings",
            # Pin the RFC-004 defaults: config load_dotenv()s a host .env that
            # may set VLM_FALLBACK=true, which would route the RFC-048
            # post-validate fallback to the VLM and break this test's hermeticity.
            dataclasses.replace(
                indexer_mod.settings, surya_fallback_enabled=False, vlm_fallback=False
            ),
        ),
    ):
        await client._convert_to_tree(
            state, str(img), img.name, ".png", None, None, script_context=latin_ctx
        )

    assert [c[0] for c in calls] == ["ensure_tessdata", "image_to_markdown"], (
        f"ensure_tessdata must precede the OCR call exactly once each; got {calls}"
    )
    assert calls[0][1] == detect_ocr_langs(img.name)
    assert calls[1][1] == provisioned[0], (
        "image_to_markdown must OCR with the langs ensure_tessdata actually provisioned"
    )
    assert seen_md and _IMAGE_OCR_MD.strip() in seen_md[0]

    egress = [
        name
        for name, double in (
            ("vlm_extract_markdown", vlm_double),
            ("_llm_with_retry", llm_double),
            ("_surya_image_ocr", surya_double),
        )
        if double.called
    ]
    assert not egress, f"HR3: image OCR route must make no model call, but called {egress}"


# ═════════════════════════════════════════════════════════════════════════
# formats.py: html_to_markdown_with_images image-describe resilience
# ═════════════════════════════════════════════════════════════════════════

# Covers the OpenAI vision call's error handling inside `_describe`:
#   - RateLimitError / APIConnectionError -> retry once after backoff
#   - non-OpenAI exceptions (TypeError etc.) propagate, are NOT swallowed
# No MinIO/Redis/network required: get_openai_client is monkeypatched.


def _counter_value(error_type: str) -> float:
    return IMAGE_DESCRIBE_FAILURES.labels(error_type=error_type)._value.get()


def _fake_request() -> httpx.Request:
    return httpx.Request("POST", "https://api.openai.com/v1/chat/completions")


def _fake_response(status_code: int = 429) -> httpx.Response:
    return httpx.Response(status_code, request=_fake_request())


def _make_client(create_mock: AsyncMock) -> SimpleNamespace:
    """Build a fake openai client shaped like client.chat.completions.create."""
    return SimpleNamespace(chat=SimpleNamespace(completions=SimpleNamespace(create=create_mock)))


def _success_response(text: str = "a picture") -> SimpleNamespace:
    return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=text))])


def _write_html(tmp_path, img_src: str = "https://example.com/pic.png") -> str:
    html_path = tmp_path / "doc.html"
    html_path.write_text(f'<html><body><img src="{img_src}"></body></html>', encoding="utf-8")
    return str(html_path)


async def test_openai_transient_errors_retry_once_then_succeed(tmp_path, monkeypatch):
    """(a) RateLimitError and APIConnectionError each retry once after a 2s
    backoff and then succeed; no "[Image: image]" fallback, no failure counter
    bump on eventual success."""
    cases = [
        (
            openai.RateLimitError("rate limited", response=_fake_response(429), body=None),
            "a cat photo",
        ),
        (
            openai.APIConnectionError(message="connection failed", request=_fake_request()),
            "a dog photo",
        ),
    ]
    failures = []
    for exc, caption in cases:
        create_mock = AsyncMock(side_effect=[exc, _success_response(caption)])
        monkeypatch.setattr(
            "pageindex_mcp.client.get_openai_client", lambda c=create_mock: _make_client(c)
        )
        sleep_mock = AsyncMock()
        monkeypatch.setattr(converters_mod.formats.asyncio, "sleep", sleep_mock)

        before = _counter_value(type(exc).__name__)
        result = await converters_mod.html_to_markdown_with_images(
            _write_html(tmp_path), model="gpt-4.1"
        )

        if caption not in result:
            failures.append(f"{type(exc).__name__}: caption missing from {result!r}")
        if "[Image: image]" in result:
            failures.append(f"{type(exc).__name__}: fell back to the 'image' placeholder")
        if create_mock.await_count != 2:
            failures.append(f"{type(exc).__name__}: {create_mock.await_count} awaits, want 2")
        sleep_mock.assert_awaited_once_with(2)
        if _counter_value(type(exc).__name__) != before:
            failures.append(f"{type(exc).__name__}: failure counter moved on a success")
    assert not failures, "\n".join(failures)


async def test_non_openai_exception_propagates(tmp_path, monkeypatch):
    """(d) A non-OpenAI exception (TypeError) is NOT caught / turned into 'image'."""
    create_mock = AsyncMock(side_effect=TypeError("boom - code bug, not an API failure"))
    monkeypatch.setattr("pageindex_mcp.client.get_openai_client", lambda: _make_client(create_mock))

    with pytest.raises(TypeError, match="boom"):
        await converters_mod.html_to_markdown_with_images(_write_html(tmp_path), model="gpt-4.1")


# ═════════════════════════════════════════════════════════════════════════
# pipeline.py: pdf_markdown_converters() chain shape + AGPL metadata
# ═════════════════════════════════════════════════════════════════════════


def _get_chain(monkeypatch, primary="docling", agpl=True):
    """Build a converter chain with controlled env vars."""
    monkeypatch.setenv("PDF_CONVERTER", primary)
    monkeypatch.setenv("ALLOW_AGPL_FALLBACK", "true" if agpl else "false")
    reset_pipeline_config()
    from pageindex_mcp.converters.pipeline import pdf_markdown_converters

    return pdf_markdown_converters()


class TestConverterChainEntryMetadata:
    """pdf_markdown_converters() returns ConverterChainEntry instances carrying
    (name, fn, supports_ocr) tuple compatibility plus the is_agpl metadata the
    chain walker needs to block transient-failure fallback to AGPL converters."""

    def test_docling_primary_chain_shape_and_agpl_metadata(self, monkeypatch):
        """With PDF_CONVERTER=docling: docling is chain[0] with supports_ocr=True
        and is_agpl=False (MIT); pymupdf4llm is present with supports_ocr=False
        and is_agpl=True (AGPL-3.0); every entry unpacks as a (name, callable,
        bool) 3-tuple of len 3 for backward compatibility."""
        chain = _get_chain(monkeypatch, primary="docling", agpl=True)
        assert len(chain) > 0
        failures = []
        for entry in chain:
            if not isinstance(entry, ConverterChainEntry):
                failures.append(f"{entry!r} is {type(entry).__name__}, not ConverterChainEntry")
                continue
            if len(entry) != 3:
                failures.append(f"{entry.name}: len(entry) == {len(entry)}, want 3")
            name, fn, supports_ocr = entry  # must not raise
            if (name, fn, supports_ocr) != (entry.name, entry.fn, entry.supports_ocr):
                failures.append(f"{entry.name}: 3-tuple unpack disagrees with the fields")
            # Indexing is part of the same backward-compat sequence protocol.
            if (entry[0], entry[1], entry[2]) != (entry.name, entry.fn, entry.supports_ocr):
                failures.append(f"{entry.name}: indexing disagrees with the fields")
            try:
                entry[3]
            except IndexError:
                pass
            else:
                failures.append(f"{entry.name}: entry[3] must raise IndexError, not expose is_agpl")
            if not isinstance(name, str) or not callable(fn) or not isinstance(supports_ocr, bool):
                failures.append(f"{entry.name}: bad (str, callable, bool) shape")
            if not isinstance(entry.is_agpl, bool):
                failures.append(
                    f"{entry.name}: is_agpl is {type(entry.is_agpl).__name__}, not bool"
                )

        by_name = {e.name: e for e in chain}
        assert "docling" in by_name, f"docling missing from chain {list(by_name)}"
        assert "pymupdf4llm" in by_name, f"pymupdf4llm missing from chain {list(by_name)}"
        if by_name["docling"].is_agpl is not False:
            failures.append("docling must have is_agpl=False (MIT-licensed)")
        if by_name["pymupdf4llm"].is_agpl is not True:
            failures.append("pymupdf4llm must have is_agpl=True (AGPL-3.0-licensed)")
        if by_name["docling"].supports_ocr is not True:
            failures.append("docling must have supports_ocr=True")
        if by_name["pymupdf4llm"].supports_ocr is not False:
            failures.append("pymupdf4llm must have supports_ocr=False")
        if (chain[0].name, chain[0].supports_ocr) != ("docling", True):
            failures.append(
                f"chain[0] is {chain[0].name}/{chain[0].supports_ocr}, want docling/True"
            )
        assert not failures, "\n".join(failures)

    def test_pymupdf_primary_ordering_and_agpl_fallback_gate(self, monkeypatch):
        """HR4: with PDF_CONVERTER=pymupdf4llm it is chain[0] (supports_ocr=False);
        with ALLOW_AGPL_FALLBACK=false the AGPL converter is absent entirely."""
        chain = _get_chain(monkeypatch, primary="pymupdf4llm", agpl=True)
        assert (chain[0].name, chain[0].supports_ocr) == ("pymupdf4llm", False)

        gated = _get_chain(monkeypatch, primary="docling", agpl=False)
        assert [e.name for e in gated if e.name == "pymupdf4llm"] == [], (
            "pymupdf4llm must not appear when ALLOW_AGPL_FALLBACK=false"
        )


# ═════════════════════════════════════════════════════════════════════════
# indexer wiring: OCR gating + AGPL chain-walk decisions
# ═════════════════════════════════════════════════════════════════════════


def _indexer_source() -> str:
    import inspect

    from pageindex_mcp.client import indexer

    return inspect.getsource(indexer)


def test_indexer_gates_on_capability_flags_not_converter_name_strings():
    """Wiring: indexer.py's chain loop unpacks the 3-tuple's third element as
    ``_conv_supports_ocr``, threads it onto ``state.supports_ocr``, and uses
    that capability flag for OCR gating / extraction_route instead of matching
    the literal string "docling".  ExtractionState carries the field and it
    defaults to False.  Also wires the AGPL chain-walk decision inputs:
    ``_classify_transient_failure``, ``ConverterChainEntry.is_agpl``, the
    ``transient_blocked`` / ``structural_walk`` / ``structural_blocked``
    AGPL_FALLBACK_TOTAL reasons, the
    ``pipeline_config.agpl_structural_fallback_enabled`` gate, and the
    ``strip_unresolved_image_markers`` safety net for empty pic_results."""
    source = _indexer_source()
    required = [
        "_conv_supports_ocr",
        "state.supports_ocr = _conv_supports_ocr",
        "_classify_transient_failure",
        ".is_agpl",
        'reason="transient_blocked"',
        'reason="structural_walk"',
        'reason="structural_blocked"',
        "pipeline_config.agpl_structural_fallback_enabled",
        "strip_unresolved_image_markers",
        "not state.pic_results",
    ]
    missing = [tok for tok in required if tok not in source]
    forbidden = ['"docling" in conv_name', '"docling" in state.used_converter']
    present = [tok for tok in forbidden if tok in source]
    assert not missing, f"indexer.py is missing required wiring: {missing}"
    assert not present, f"indexer.py still string-matches the converter name: {present}"

    field_names = [f.name for f in dataclasses.fields(ExtractionState)]
    assert "supports_ocr" in field_names
    state = ExtractionState(
        result={},
        ok=False,
        reason="",
        gate_result=None,
        first_defect=TreeDefect.OK,
        route=Route.TREE,
        md_content=None,
        tmp_md_path=None,
        pic_results=[],
        used_converter=None,
        total_chars=0,
        extraction_stages_captured=[],
    )
    assert state.supports_ocr is False


def test_classify_transient_failure_partitions_transient_from_structural():
    """HR4 input: transient failures (TimeoutError, ConnectionError, OSError,
    HTTP 5xx) must be separable from structural ones (ValueError, RuntimeError,
    ImportError, HTTP 4xx), because only structural failures may justify a walk
    into an AGPL-licensed converter."""
    from pageindex_mcp.client.indexer import _classify_transient_failure

    http_5xx = Exception("gateway timeout")
    http_5xx.status_code = 504  # type: ignore[attr-defined]
    http_4xx = Exception("bad request")
    http_4xx.status_code = 400  # type: ignore[attr-defined]

    cases = [
        (TimeoutError("timed out"), True),
        (ConnectionError("refused"), True),
        (OSError("network unreachable"), True),
        (http_5xx, True),
        (ValueError("bad format"), False),
        (RuntimeError("empty output"), False),
        (ImportError("no module"), False),
        (http_4xx, False),
    ]
    failures = [
        f"{exc!r}: _classify_transient_failure -> {_classify_transient_failure(exc)}, want {want}"
        for exc, want in cases
        if _classify_transient_failure(exc) is not want
    ]
    assert not failures, "\n".join(failures)


# ═════════════════════════════════════════════════════════════════════════
# indexer chain walk: ConverterFailurePolicy.GATE_AGPL_STRUCTURAL
# ═════════════════════════════════════════════════════════════════════════

# A STRUCTURAL failure that would walk into an AGPL-licensed converter is an
# explicit, metricked, operator-gateable policy branch instead of an unnamed
# fall-through into WALK.

_GATE_TREE = {
    "structure": [
        {
            "node_id": "0001",
            "title": "Root",
            "text": "content " * 60,
            "nodes": [{"node_id": "0002", "title": "Child", "text": "child " * 60, "nodes": []}],
        }
    ]
}


def _make_gate_state():
    """Fresh ExtractionState for driving the real _convert_to_tree chain walk."""
    return ExtractionState(
        result={},
        ok=False,
        reason="",
        gate_result=None,
        first_defect=TreeDefect.OK,
        route=Route.TREE,
        md_content=None,
        tmp_md_path=None,
        pic_results=[],
        used_converter=None,
        total_chars=0,
        extraction_stages_captured=[],
    )


def _make_gate_client():
    """CustomPageIndexClient with the LLM tree-builders stubbed out.

    ``_convert_to_tree`` is exercised for real; only the two terminal
    tree-producing coroutines are replaced so no LLM/page_index call is made.
    ``_staging_key = None`` forces the local (non-remote) converter branch.
    """
    from pageindex_mcp.client import CustomPageIndexClient

    client = CustomPageIndexClient(api_key="test-key")
    client._staging_key = None
    client._run_md_to_tree = AsyncMock(return_value=dict(_GATE_TREE))
    client._run_page_index_retrying = AsyncMock(return_value=dict(_GATE_TREE))
    return client


def _agpl_metric(reason: str) -> float:
    from pageindex_mcp.metrics import AGPL_FALLBACK_TOTAL

    return AGPL_FALLBACK_TOTAL.labels(reason=reason)._value.get()


def _pdf_extract_fallbacks() -> float:
    from pageindex_mcp.metrics import PDF_EXTRACT_FALLBACKS

    return PDF_EXTRACT_FALLBACKS._value.get()


async def _run_chain(chain, *, structural_fallback_enabled=True):
    """Drive the real ``_convert_to_tree`` chain walk over *chain*.

    Returns ``(client, state)`` so callers can assert on which converter won,
    whether the legacy page_index fallback fired, and which metric moved.
    """
    from pageindex_mcp.client import indexer as indexer_mod

    client = _make_gate_client()
    state = _make_gate_state()
    patched_cfg = dataclasses.replace(
        indexer_mod.pipeline_config,
        agpl_structural_fallback_enabled=structural_fallback_enabled,
    )
    with (
        patch.object(indexer_mod, "pdf_markdown_converters", lambda: list(chain)),
        patch.object(indexer_mod, "pipeline_config", patched_cfg),
    ):
        await client._convert_to_tree(
            state,
            "/nonexistent/zone-gate-fixture.pdf",
            "zone-gate-fixture.pdf",
            ".pdf",
            "latin",
            None,
        )
    return client, state


def _gate_chain(first_exc):
    """non-AGPL primary that raises *first_exc* -> AGPL fallback that succeeds."""
    primary = MagicMock(side_effect=first_exc)
    agpl = MagicMock(return_value=("# recovered markdown", [], []))
    chain = [
        ConverterChainEntry(name="docling", fn=primary, supports_ocr=True, is_agpl=False),
        ConverterChainEntry(name="pymupdf4llm", fn=agpl, supports_ocr=False, is_agpl=True),
    ]
    return chain, primary, agpl


class TestGateAgplStructuralPolicy:
    """GATE_AGPL_STRUCTURAL: structural failure walking into an AGPL converter."""

    def test_enum_is_exhaustively_classified_and_dispatched_in_indexer(self):
        """The enum carries GATE_AGPL_STRUCTURAL = 'gate_agpl_structural', and
        every ConverterFailurePolicy member is both assigned and dispatched on
        in the indexer chain walker -- no member may exist without a branch
        that acts on it."""
        from pageindex_mcp.converters.pipeline import ConverterFailurePolicy

        assert ConverterFailurePolicy.GATE_AGPL_STRUCTURAL.value == "gate_agpl_structural"
        assert ConverterFailurePolicy("gate_agpl_structural") is (
            ConverterFailurePolicy.GATE_AGPL_STRUCTURAL
        )

        source = _indexer_source()
        failures = []
        for member in ConverterFailurePolicy:
            ref = f"ConverterFailurePolicy.{member.name}"
            if source.count(ref) < 2:
                failures.append(
                    f"{ref} must be both classified and handled in indexer.py "
                    f"(found {source.count(ref)} reference(s))"
                )
            # WALK is the documented implicit tail branch (no `is WALK` compare);
            # every other member must be dispatched with an identity check.
            if member is not ConverterFailurePolicy.WALK:
                if f"_failure_policy is ConverterFailurePolicy.{member.name}" not in source:
                    failures.append(f"indexer.py has no dispatch branch for {ref}")
        assert not failures, "\n".join(failures)

    @pytest.mark.asyncio
    async def test_structural_failure_walk_to_agpl_is_gated_by_config(self):
        """Contract: a structural failure (ValueError) on a non-AGPL converter
        walks into the AGPL converter when the gate is enabled (default),
        counted as AGPL_FALLBACK_TOTAL(reason='structural_walk'); with
        AGPL_STRUCTURAL_FALLBACK_ENABLED=false the AGPL converter is never
        invoked, the document falls to the legacy page_index path, and the
        block is counted as reason='structural_blocked'.

        INDEX-01-C2: the .pdf -> ``_run_page_index`` last-resort fallback fires
        ONLY when the markdown route raised (here: every converter in the chain
        failed), and that fallback invocation is OBSERVABLE -- it increments
        ``PDF_EXTRACT_FALLBACKS`` and emits the
        ``pdf_conversion_outcome=all_converters_failed_legacy_fallback``
        decision record.  On the success half neither the fallback nor the
        counter moves."""
        chain, primary, agpl = _gate_chain(ValueError("unparseable PDF structure"))
        before_walk = _agpl_metric("structural_walk")
        before_blocked = _agpl_metric("structural_blocked")
        before_fallbacks = _pdf_extract_fallbacks()

        client, state = await _run_chain(chain, structural_fallback_enabled=True)

        primary.assert_called_once()
        agpl.assert_called_once()
        assert state.used_converter == "pymupdf4llm", (
            "structural failure must walk to the AGPL converter when enabled"
        )
        assert state.md_content == "# recovered markdown"
        client._run_page_index_retrying.assert_not_called()
        assert _agpl_metric("structural_walk") == before_walk + 1
        assert _agpl_metric("structural_blocked") == before_blocked
        assert _pdf_extract_fallbacks() == before_fallbacks, (
            "INDEX-01-C2: no legacy page_index fallback when a converter succeeded"
        )

        chain, primary, agpl = _gate_chain(ValueError("unparseable PDF structure"))
        from pageindex_mcp.client import indexer as _indexer_mod

        with patch.object(_indexer_mod, "decision", MagicMock()) as decision_mock:
            client, state = await _run_chain(chain, structural_fallback_enabled=False)

        primary.assert_called_once()
        agpl.assert_not_called()
        assert state.used_converter is None
        assert state.md_content is None
        client._run_page_index_retrying.assert_called_once()
        assert _agpl_metric("structural_blocked") == before_blocked + 1
        assert _agpl_metric("structural_walk") == before_walk + 1
        # INDEX-01-C2: the fallback must be observable, not silent.
        assert _pdf_extract_fallbacks() == before_fallbacks + 1
        assert [
            call.kwargs.get("choice")
            for call in decision_mock.call_args_list
            if call.kwargs.get("event") == "pdf_conversion_outcome"
        ] == ["all_converters_failed_legacy_fallback"]

    @pytest.mark.asyncio
    async def test_transient_to_agpl_still_blocks_unchanged(self):
        """HR4 regression: the pre-existing BLOCK_AGPL branch is untouched by
        the GATE_AGPL_STRUCTURAL branch -- a TRANSIENT failure into an AGPL
        converter still blocks the walk (the AGPL converter is never invoked)
        and still counts as reason='transient_blocked', with the structural
        gate enabled.

        Retries are disabled here so the BLOCK_AGPL branch is reached on the
        first failure; see ``test_transient_retry_does_not_reenter_same_converter``
        for the RETRY branch's separate (defective) behavior.
        """
        from pageindex_mcp.client import indexer as indexer_mod

        chain, primary, agpl = _gate_chain(TimeoutError("docling timed out"))
        before_transient = _agpl_metric("transient_blocked")
        before_walk = _agpl_metric("structural_walk")
        before_blocked = _agpl_metric("structural_blocked")

        with patch.object(indexer_mod, "CONVERTER_TRANSIENT_RETRY_COUNT", 0):
            client, state = await _run_chain(chain, structural_fallback_enabled=True)

        primary.assert_called_once()
        agpl.assert_not_called()
        assert state.used_converter is None
        client._run_page_index_retrying.assert_called_once()
        assert _agpl_metric("transient_blocked") == before_transient + 1
        # The structural counters must not move for a transient failure.
        assert _agpl_metric("structural_walk") == before_walk
        assert _agpl_metric("structural_blocked") == before_blocked

    @pytest.mark.xfail(
        strict=False,
        reason=(
            "DEFECT: indexer.py's RETRY branch `continue`s a plain "
            "`for idx, entry in enumerate(chain)` loop. The comment claims it "
            "'rewinds idx so the for-loop re-enters this entry', but nothing "
            "rewinds -- the continue advances to the NEXT converter. A transient "
            "failure therefore walks straight into the AGPL converter on the "
            "first attempt, bypassing BLOCK_AGPL (HR4) entirely, and the "
            "configured retry never happens."
        ),
    )
    @pytest.mark.asyncio
    async def test_transient_retry_does_not_reenter_same_converter(self):
        """RETRY must re-invoke the SAME converter, not advance the chain --
        otherwise CONVERTER_TRANSIENT_RETRY_COUNT>0 silently defeats the HR4
        AGPL block for transient failures."""
        from pageindex_mcp.client import indexer as indexer_mod

        chain, primary, agpl = _gate_chain(TimeoutError("docling timed out"))
        with patch.object(indexer_mod, "CONVERTER_TRANSIENT_RETRY_COUNT", 2):
            client, state = await _run_chain(chain, structural_fallback_enabled=True)

        assert primary.call_count == 3, "RETRY must re-invoke the same converter"
        agpl.assert_not_called(), "transient failure must never reach the AGPL converter"


# ═════════════════════════════════════════════════════════════════════════
# pictures.py: _bbox_to_fitz_rect / splice_figure_markers
# ═════════════════════════════════════════════════════════════════════════


class _FakeRect:
    def __init__(self, x0, y0, x1, y1):
        self.x0, self.y0, self.x1, self.y1 = x0, y0, x1, y1


def test_bbox_to_fitz_rect_handles_both_coord_origins():
    """RFC-015 D6: ``_bbox_to_fitz_rect`` passes TOPLEFT-origin bboxes through
    unchanged and flips BOTTOMLEFT-origin ones against the page height
    (top = 800-700 = 100, bottom = 800-600 = 200 -> sorted y (100, 200))."""
    fitz = types.SimpleNamespace(Rect=_FakeRect)

    topleft = types.SimpleNamespace(l=10, t=20, r=110, b=120, coord_origin=None)
    rect = _bbox_to_fitz_rect(topleft, 800.0, fitz)
    assert (rect.x0, rect.y0, rect.x1, rect.y1) == (10, 20, 110, 120)

    bottomleft = types.SimpleNamespace(
        l=10, t=700, r=110, b=600, coord_origin=types.SimpleNamespace(name="BOTTOMLEFT")
    )
    rect = _bbox_to_fitz_rect(bottomleft, 800.0, fitz)
    assert (rect.y0, rect.y1) == (100, 200)


def test_splice_figure_markers_replaces_markers_and_appends_chart_text():
    """RFC-015 D6 / audit findings 4+7+12: ``splice_figure_markers`` replaces
    ``<!-- image -->`` markers with ``[Figure: fig-N]`` refs from a DENSE
    ordinal-keyed list and appends recovered chart text as a blockquote; with
    no picture results the markdown is returned unchanged.  Accepts both the
    dict shape and the ``PictureResult`` dataclass."""
    from pageindex_mcp.converters import PictureResult

    pr = {"ocr_text": "Revenue 2024 42%", "png_bytes": b"png", "page": 1, "bbox": {}}
    out = splice_figure_markers("Intro\n\n<!-- image -->\n\nOutro", [pr])
    assert "[Figure: fig-0]" in out
    assert "> [Chart text]: Revenue 2024 42%" in out
    assert "<!-- image -->" not in out

    assert splice_figure_markers("<!-- image -->", []) == "<!-- image -->"

    dataclass_out = splice_figure_markers(
        "Text <!-- image --> more",
        [PictureResult(ocr_text="Chart data", page=0, bbox={"l": 0, "t": 0, "r": 100, "b": 100})],
    )
    assert "[Figure: fig-0]" in dataclass_out


# ═════════════════════════════════════════════════════════════════════════
# pictures.py: _recover_picture_results escalation gate + crash isolation
# ═════════════════════════════════════════════════════════════════════════


class TestRecoverPictureResults:
    """RFC-015 D6 / audit finding 6: ``_recover_picture_results`` gates the
    first-party AGPL ``fitz`` import (via ``_recover_picture_text``) behind the
    per-picture OCR-escalation config, and NEVER mutates the markdown -- the
    figure splice happens only in client.index()'s flat branch."""

    def test_escalation_gate_controls_whether_recovery_runs(self, monkeypatch):
        md = "Intro\n\n<!-- image -->\n\nOutro"
        bbox = types.SimpleNamespace(l=0, t=10, r=100, b=110, coord_origin=None)
        pictures = [{"page": 1, "bbox": bbox}]

        # Zone-5 config layering: the gate reads the pipeline_config singleton
        # at call time, not the frozen module-level alias.
        monkeypatch.setattr(
            converters.pictures,
            "pipeline_config",
            dataclasses.replace(
                converters.pictures.pipeline_config, ocr_escalation_per_picture=False
            ),
        )
        with (
            mock.patch.object(
                converters.pictures, "_collect_picture_regions", return_value=pictures
            ) as mock_collect,
            mock.patch.object(converters.pictures, "_recover_picture_text") as mock_recover,
        ):
            pics = converters._recover_picture_results(md, object(), "dummy.pdf")
        mock_collect.assert_not_called()
        mock_recover.assert_not_called()
        assert pics == []

        monkeypatch.setattr(
            converters.pictures,
            "pipeline_config",
            dataclasses.replace(
                converters.pictures.pipeline_config, ocr_escalation_per_picture=True
            ),
        )
        monkeypatch.setattr(converters.pictures, "_OCR_ESCALATION_PER_PICTURE", True)
        pr = {
            "ocr_text": "Revenue 2024 recovered chart text",
            "png_bytes": b"fake",
            "page": 1,
            "bbox": {},
        }
        with (
            mock.patch.object(
                converters.pictures, "_collect_picture_regions", return_value=pictures
            ),
            mock.patch.object(converters.pictures, "detect_ocr_langs", return_value=["eng"]),
            mock.patch.object(
                converters.pictures, "ensure_tessdata", side_effect=lambda langs: langs
            ),
            mock.patch.object(
                converters.pictures, "_recover_picture_text", return_value=({0: pr}, {})
            ) as mock_recover,
        ):
            pics = converters._recover_picture_results(md, object(), "dummy.pdf")
        assert mock_recover.call_count >= 1
        assert pics == [pr]

    def test_recover_picture_results_returns_empty_on_total_failure(self, monkeypatch):
        """When ``_recover_picture_text`` itself raises (e.g. the PDF cannot be
        opened at all), the outer except in ``_recover_picture_results`` still
        returns an empty list rather than propagating."""
        monkeypatch.setattr(converters.pictures, "_OCR_ESCALATION_PER_PICTURE", True)
        monkeypatch.setattr(
            converters.pictures,
            "_collect_picture_regions",
            lambda document: [_region(0, 0, 30, 30)],
        )

        def _boom(pdf_path, regions, langs, md=""):
            raise RuntimeError("pdf could not be opened")

        monkeypatch.setattr(converters.pictures, "_recover_picture_text", _boom)

        result = _recover_picture_results(
            "some heading\n\n<!-- image -->\n\nmore text", document=object(), pdf_path="/fake.pdf"
        )
        assert result == []

    def test_per_region_crop_failure_is_isolated_not_propagated(self, monkeypatch):
        """RFC-024 D2: a degenerate region whose ``get_pixmap`` raises is tagged
        ``crop_error`` rather than aborting the document; when every region
        fails the call still returns gracefully."""
        fake_fitz, _page = _make_fake_fitz(600.0, 800.0, pixmap_fails_for_l={100, 200, 300})
        monkeypatch.setattr(converters.pictures, "_tesseract_ocr_image", lambda path, langs: "")
        regions = [_region(100, 0, 130, 30), _region(200, 0, 230, 30), _region(300, 0, 330, 30)]

        with patch.dict(sys.modules, {"fitz": fake_fitz}):
            result, skip_reasons = _recover_picture_text("/fake.pdf", regions, ["eng"])

        assert result == {}
        assert set(skip_reasons.values()) == {"crop_error"}
        assert len(skip_reasons) == 3


# ═════════════════════════════════════════════════════════════════════════
# garble detection (helpers.garble via the converter profiles)
# ═════════════════════════════════════════════════════════════════════════


def test_tree_and_flat_garble_detection_prongs():
    """D3A/D3B: the PUA-ratio (>3%, font/CMap mojibake) and digit-ratio (>60%
    on a blob >500 chars) prongs fire on both the tree-bulk profile and the
    flat-markdown profile."""
    pua_text_a = "" * 5 + "a" * 90
    pua_text_b = "" * 5 + "b" * 90
    digit_text = "1651001429 " * 80  # ~880 chars, >60% digits

    pua_nodes = [
        {
            "title": "X",
            "text": pua_text_a,
            "nodes": [{"title": "Y", "text": pua_text_b, "nodes": []}],
        }
    ]
    digit_nodes = [
        {
            "title": "A",
            "text": digit_text,
            "nodes": [{"title": "B", "text": "some text", "nodes": []}],
        }
    ]
    cases = [
        ("tree/pua", _tree_garble(pua_nodes)),
        ("tree/digit", _tree_garble(digit_nodes)),
        ("flat/pua", _flat_garble(pua_text_a + pua_text_b)),
        ("flat/digit", _flat_garble(digit_text)),
    ]
    not_flagged = [label for label, got in cases if got is not True]
    assert not not_flagged, f"garble prongs failed to fire for: {not_flagged}"


def test_image_markers_exempt_from_garble_but_real_repetition_is_not():
    """Design Property 4 (D3): a blob consisting solely of ``<!-- ... -->`` HTML
    comment markers is never flagged garbled (they are structural markers, not
    mojibake, and a scanned PDF's markdown is 100% single-token repetition
    pre-D3); genuine repeated non-comment tokens above the 30% threshold
    still are."""
    assert (
        check_garble("\n\n".join([_IMAGE_MARKER] * 45), expected_script=None, profile=BULK_PROFILE)
        is False
    )
    assert check_garble(_GARBLED_TEXT, expected_script=None, profile=BULK_PROFILE) is True


def test_short_post_retry_text_is_not_force_flagged_garbled(monkeypatch):
    """D2 + Zone-7 fix: a prior GARBLING / NODE_GARBLING defect no longer
    force-flags clean short text -- the real prongs run first and none fire on
    this clean policy sentence.  Unrelated defects get the same normal
    evaluation, and GARBLE_SHORT_TEXT_DEFAULT=false (rollback) does not change
    the answer for clean text either."""
    cases = [
        (True, TreeDefect.GARBLING),
        (True, TreeDefect.NODE_GARBLING),
        (True, TreeDefect.NODE_COUNT_LOW),
        (False, TreeDefect.GARBLING),
    ]
    failures = []
    for short_default, defect in cases:
        monkeypatch.setattr(helpers.garble, "_GARBLE_SHORT_TEXT_DEFAULT", short_default)
        got = check_garble(
            _SHORT_CLEAN_TEXT,
            expected_script=None,
            profile=FLAT_MARKDOWN_PROFILE,
            original_defect=defect,
        )
        if got is not False:
            failures.append(
                f"short_default={short_default}, defect={defect.value}: got {got}, want False"
            )
    assert not failures, "\n".join(failures)


# ═════════════════════════════════════════════════════════════════════════
# pictures.py: _tesseract_ocr_image exception-handling contract (Zone-8)
# ═════════════════════════════════════════════════════════════════════════


class TestTesseractOcrFailureContract:
    """Zone-8: ``_tesseract_ocr_image`` increments TESSERACT_OCR_FAILURE_TOTAL
    with the exception class name and returns '' for the handled failure
    modes -- and does NOT catch arbitrary exceptions like KeyboardInterrupt."""

    def test_handled_exceptions_increment_metric_and_return_empty(self, monkeypatch):
        from pageindex_mcp.converters.pictures import _tesseract_ocr_image

        monkeypatch.setattr("shutil.which", lambda name: "/usr/bin/tesseract")
        cases = [
            (subprocess.TimeoutExpired(cmd="tesseract", timeout=60), "TimeoutExpired"),
            (subprocess.SubprocessError("boom"), "SubprocessError"),
            (FileNotFoundError("tesseract not found"), "FileNotFoundError"),
            (OSError("disk error"), "OSError"),
        ]
        failures = []
        for exc, reason in cases:
            with (
                patch("pageindex_mcp.converters.pictures.subprocess.run", side_effect=exc),
                patch("pageindex_mcp.converters.pictures.TESSERACT_OCR_FAILURE_TOTAL") as metric,
            ):
                result = _tesseract_ocr_image("/fake.png", ["eng"])
            if result != "":
                failures.append(f"{reason}: returned {result!r}, want ''")
            try:
                metric.labels.assert_called_once_with(reason=reason)
                metric.labels.return_value.inc.assert_called_once()
            except AssertionError as err:
                failures.append(f"{reason}: metric not incremented ({err})")
        assert not failures, "\n".join(failures)

    def test_keyboard_interrupt_not_caught(self, monkeypatch):
        """KeyboardInterrupt must NOT be caught -- it must propagate."""
        from pageindex_mcp.converters.pictures import _tesseract_ocr_image

        monkeypatch.setattr("shutil.which", lambda name: "/usr/bin/tesseract")
        with (
            patch(
                "pageindex_mcp.converters.pictures.subprocess.run", side_effect=KeyboardInterrupt
            ),
            pytest.raises(KeyboardInterrupt),
        ):
            _tesseract_ocr_image("/fake.png", ["eng"])


# ═════════════════════════════════════════════════════════════════════════
# pictures.py: _text_layer_has_content (page-level and region-scoped)
# ═════════════════════════════════════════════════════════════════════════


def _page(text: str):
    return types.SimpleNamespace(get_text=lambda mode="text": text)


def test_text_layer_has_content_page_level_requires_clean_and_long_enough():
    """Design Property 1: ``_text_layer_has_content`` returns False for text
    that is either too short or flagged garbled (thin mojibake left by the PDF
    creator must not be treated as real content), and True only when both
    checks pass."""
    from pageindex_mcp.helpers import GarbleReport

    assert _text_layer_has_content(_page(_GARBLED_TEXT)) is False
    assert _text_layer_has_content(_page(_CLEAN_TEXT)) is True

    # Explicitly drive the garble branch independent of the heuristics.
    garbled = GarbleReport(is_garbled=True, fired_prongs=frozenset({"test"}))
    with patch("pageindex_mcp.converters.pictures.detect_garble", return_value=garbled):
        assert _text_layer_has_content(_page("A" * 100)) is False


def test_text_layer_has_content_region_scoped_min_chars_threshold():
    """D1: the region-scoped (clipped) read drives the same check -- header
    text living outside the region's own bbox yields an empty clipped read, so
    the region has NO text of its own.  ``_PICTURE_OCR_MIN_CHARS`` is 20 and
    the length check is strict (``len(text) <= _PICTURE_OCR_MIN_CHARS`` fails
    at exactly 20), so the smallest length that clears the threshold is 21."""
    rect = types.SimpleNamespace()
    cases = [("", False), ("x" * 19, False), ("x" * 20, False), ("The quick brown foxes", True)]
    failures = []
    for text, want in cases:
        page = types.SimpleNamespace(get_text=lambda mode, clip=None, t=text: t)
        got = _text_layer_has_content(page, region_rect=rect)
        if got is not want:
            failures.append(f"len={len(text)}: got {got}, want {want}")
    assert not failures, "\n".join(failures)


# ═════════════════════════════════════════════════════════════════════════
# pictures.py: region-aware coverage exemption + full-page region cap
# ═════════════════════════════════════════════════════════════════════════


class TestRegionAwareExemptionIntegration:
    """Region-scoped text-layer check wired into ``_recover_picture_text``."""

    def test_exemption_fires_only_when_the_regions_own_bbox_is_textless(self, monkeypatch):
        """A page with header/footer text (where the page-level check would see
        content and skip) but whose picture bbox has none fires the
        region-aware exemption and OCR proceeds.  Conversely, when the region's
        own bbox carries real text the exemption must NOT fire -- the
        region-scoped check must not become permissive in the other direction
        (edge case from the design doc).  RFC-029 D5a: skipped regions still
        surface in ``recovered`` carrying ``png_bytes`` + ``skipped_reason``,
        but ``ocr_text`` MUST be absent -- proving Tesseract was not run."""
        monkeypatch.setattr(converters.pictures, "_COVERAGE_EXEMPT_NO_TEXT_LAYER", True)
        monkeypatch.setattr(
            converters.pictures, "_tesseract_ocr_image", lambda png, langs: _long_text()
        )

        _install_fake_fitz(monkeypatch, page_text=_long_text(60), clip_text="")
        recovered, skip_reasons = _recover_picture_text("dummy.pdf", [_region()], ["eng"])
        assert skip_reasons.get(0) != "page_coverage"
        assert recovered[0]["ocr_text"] == _long_text()

        _install_fake_fitz(monkeypatch, page_text="", clip_text=_long_text(60))
        recovered, skip_reasons = _recover_picture_text("dummy.pdf", [_region()], ["eng"])
        assert skip_reasons.get(0) == "page_coverage"
        assert "ocr_text" not in recovered.get(0, {})
        assert recovered.get(0, {}).get("skipped_reason") == "page_coverage"

    def test_regions_past_cap_skipped_with_page_coverage(self, monkeypatch):
        """``MAX_FULLPAGE_PICTURE_OCR_REGIONS`` is a per-document boundary
        defaulting to 50.  With the cap set to 2 and 3 qualifying full-page
        regions, the first 2 get the exemption and OCR fires; the 3rd is
        skipped with "page_coverage" and a logged warning, not silently
        exempted (RFC-029 D5a: retained with ``png_bytes`` + ``skipped_reason``
        but WITHOUT ``ocr_text``)."""
        assert converters._MAX_FULLPAGE_PICTURE_OCR_REGIONS == 50

        monkeypatch.setattr(converters.pictures, "_COVERAGE_EXEMPT_NO_TEXT_LAYER", True)
        monkeypatch.setattr(converters.pictures, "_MAX_FULLPAGE_PICTURE_OCR_REGIONS", 2)
        monkeypatch.setattr(
            converters.pictures,
            "_GATE_CONFIG",
            PictureGateConfig(
                coverage_exempt_no_text_layer=True,
                max_fullpage_picture_ocr_regions=2,
            ),
        )
        _install_fake_fitz(monkeypatch, page_text=_long_text(60), clip_text="")
        monkeypatch.setattr(
            converters.pictures, "_tesseract_ocr_image", lambda png, langs: _long_text()
        )

        regions = [_region() for _ in range(3)]
        recovered, skip_reasons = _recover_picture_text("dummy.pdf", regions, ["eng"])

        assert skip_reasons.get(0) != "page_coverage"
        assert skip_reasons.get(1) != "page_coverage"
        assert skip_reasons.get(2) == "page_coverage"
        assert recovered[0]["ocr_text"] == _long_text()
        assert recovered[1]["ocr_text"] == _long_text()
        assert "ocr_text" not in recovered.get(2, {})
        assert recovered.get(2, {}).get("skipped_reason") == "page_coverage"


def test_decorative_flag_has_no_rotation_gate(monkeypatch):
    """The rotation gate is removed: empty OCR sets the decorative
    ``ocr_min_chars`` skip reason even when rotation != 0 (previously it only
    fired at rotation == 0), while non-empty OCR on a rotated page sets no
    skip reason at all."""
    region = _region(0, 0, 30, 30)

    fake_fitz, _page = _make_fake_fitz(600.0, 800.0, initial_rotation=180)
    monkeypatch.setattr(converters.pictures, "_tesseract_ocr_image", lambda path, langs: "")
    with patch.dict("sys.modules", {"fitz": fake_fitz}):
        result, _skip = _recover_picture_text("/fake.pdf", [region], ["eng"])
    assert result[0].get("skipped_reason") == "ocr_min_chars"

    fake_fitz, _page = _make_fake_fitz(600.0, 800.0, initial_rotation=90)
    monkeypatch.setattr(
        converters.pictures,
        "_tesseract_ocr_image",
        lambda path, langs: "Recovered chart text with enough characters",
    )
    with patch.dict("sys.modules", {"fitz": fake_fitz}):
        result, _skip = _recover_picture_text("/fake.pdf", [region], ["eng"])
    assert "skipped_reason" not in result[0]


def test_decorative_icon_size_filter(monkeypatch):
    """Design Property 3: a PictureItem region whose bbox width AND height are
    both below DECORATIVE_ICON_MIN_DIM_PT skips crop+OCR entirely and is tagged
    ``skip_reasons[i] == "decorative_icon"``; a region above the threshold
    proceeds to OCR."""
    monkeypatch.setattr(converters.pictures, "_DECORATIVE_ICON_MIN_DIM_PT", 20.0)

    fake_fitz, _page = _make_fake_fitz(600.0, 800.0)

    def _fail_if_called(*_a, **_k):
        raise AssertionError("tesseract must not run for sub-icon regions")

    monkeypatch.setattr(converters.pictures, "_tesseract_ocr_image", _fail_if_called)
    with patch.dict("sys.modules", {"fitz": fake_fitz}):
        result, skip_reasons = _recover_picture_text("/fake.pdf", [_region(0, 0, 15, 12)], ["eng"])
    assert result == {}
    assert skip_reasons[0] == "decorative_icon"

    fake_fitz, _page = _make_fake_fitz(600.0, 800.0)
    monkeypatch.setattr(
        converters.pictures,
        "_tesseract_ocr_image",
        lambda path, langs: "Chart text with enough characters to pass the gate",
    )
    with patch.dict("sys.modules", {"fitz": fake_fitz}):
        result, skip_reasons = _recover_picture_text("/fake.pdf", [_region(0, 0, 30, 30)], ["eng"])
    assert 0 not in skip_reasons
    assert result[0]["ocr_text"]


class TestRecoverPictureTextClipCapture:
    """RFC-024 D1: a region's clipped text layer is captured directly (no
    Tesseract) when it is not already present in the markdown, and skipped
    when it is -- guarded by a whitespace/reflow-robust containment check."""

    def _clip_fitz(self, clip_text):
        fake = types.ModuleType("fitz")
        fake.Rect = lambda *a: types.SimpleNamespace(
            l=a[0], t=a[1], r=a[2], b=a[3], width=a[2] - a[0], height=a[3] - a[1]
        )

        class _FakePage:
            def __init__(self):
                self.rect = types.SimpleNamespace(height=800.0, width=600.0)
                self.rotation = 0

            def get_text(self, mode="text", *, clip=None):
                return clip_text

            def set_rotation(self, value):
                self.rotation = value

            def get_pixmap(self, *, clip=None, dpi=300):
                raise AssertionError("tesseract crop path must not run when clip_text is captured")

        page = _FakePage()

        class _FakeDoc:
            page_count = 1

            def __getitem__(self, idx):
                return page

            def close(self):
                pass

        fake.open = lambda path: _FakeDoc()
        return fake

    def test_clip_text_captured_unless_already_exported(self, monkeypatch):
        clip_text = "Revenue grew 42% year over year across all regions"
        md = "# Report\n\nSome unrelated heading content.\n\n<!-- image -->"

        def _fail_if_called(path, langs):
            raise AssertionError("tesseract must not run for captured clip_text")

        monkeypatch.setattr(converters.pictures, "_tesseract_ocr_image", _fail_if_called)
        with patch.dict(sys.modules, {"fitz": self._clip_fitz(clip_text)}):
            result, skip_reasons = _recover_picture_text(
                "/fake.pdf", [_region(0, 0, 100, 40)], ["eng"], md=md
            )
        assert 0 not in skip_reasons
        assert result[0]["ocr_text"] == clip_text

        already = "The quarterly revenue increased significantly this year"
        monkeypatch.setattr(
            converters.pictures, "_tesseract_ocr_image", lambda path, langs: "should not matter"
        )
        with patch.dict(sys.modules, {"fitz": self._clip_fitz(already)}):
            result, skip_reasons = _recover_picture_text(
                "/fake.pdf",
                [_region(0, 0, 100, 40)],
                ["eng"],
                md=f"# Report\n\n{already}\n\n<!-- image -->",
            )
        assert skip_reasons[0] == "clip_text_already_exported"
        assert 0 not in result

    def test_containment_helper_matches_reflowed_text_and_rejects_unrelated(self):
        """``_clip_text_contained`` is whitespace/reflow-robust and returns
        False for genuinely unrelated content."""
        md_norm = _normalize_for_containment(
            "revenue grew 42% year-over-year across all business units"
        )
        assert _clip_text_contained("Revenue   grew\n42%  year-over-year", md_norm) is True

        unrelated_md = _normalize_for_containment(
            "This markdown body talks about something else entirely."
        )
        assert (
            _clip_text_contained("Completely unrelated chart label content here", unrelated_md)
            is False
        )


# ═════════════════════════════════════════════════════════════════════════
# _document_level_text_fallback
# ═════════════════════════════════════════════════════════════════════════


def _fake_pdfium_module(page_texts):
    fake = types.ModuleType("pypdfium2")

    class _FakeTextPage:
        def __init__(self, text):
            self._text = text

        def get_text_range(self):
            return self._text

    class _FakePage:
        def __init__(self, text):
            self._text = text

        def get_textpage(self):
            return _FakeTextPage(self._text)

    class _FakeDoc:
        def __init__(self, texts):
            self._pages = [_FakePage(t) for t in texts]

        def __iter__(self):
            return iter(self._pages)

        def close(self):
            pass

    fake.PdfDocument = lambda path: _FakeDoc(page_texts)
    return fake


class TestHeadingOnlyFallbackTrigger:
    """Chars-per-heading secondary trigger for ``_document_level_text_fallback``
    (heading-only trees where structure survived but body prose did not)."""

    def test_chars_per_heading_floor_gates_the_fallback(self, monkeypatch):
        """6 headings with ~40 chars of body text between them (~7
        chars/heading, well under the 50-char floor) trigger the fallback even
        though total_chars clears the absolute 100-char floor; 2 headings with
        well over 50 chars/heading of prose must NOT fire it -- pdfium is not
        even invoked and the markdown comes back unchanged."""
        md = "\n\n".join(f"# Heading {i}\n\nshort" for i in range(6))
        assert (
            len(md.replace(converters._IMAGE_MARKER, "")) >= converters._DOC_TEXT_FALLBACK_MIN_CHARS
        )

        recovered = "Recovered whole-document prose that clears the garble floor easily."
        monkeypatch.setitem(sys.modules, "pypdfium2", _fake_pdfium_module([recovered]))
        result = _document_level_text_fallback(md, "/fake.pdf")
        assert result != md
        assert recovered in result

        rich = "# Heading 1\n\n" + ("word " * 40) + "\n\n# Heading 2\n\n" + ("word " * 40)
        fake_pdfium = types.ModuleType("pypdfium2")
        fake_pdfium.PdfDocument = lambda path: (_ for _ in ()).throw(
            AssertionError("pdfium should not be invoked when chars/heading clears the floor")
        )
        monkeypatch.setitem(sys.modules, "pypdfium2", fake_pdfium)
        assert _document_level_text_fallback(rich, "/fake.pdf") == rich

    def test_fallback_leaves_markdown_unchanged_on_garble_or_open_failure(self, monkeypatch):
        """RFC-024 D1 risk mitigation: a scanned page's thin mojibake text layer
        must never be appended (HR5), and a pdfium open failure must degrade
        gracefully -- both return the markdown unchanged."""
        from pageindex_mcp.helpers import GarbleReport

        md = "<!-- image -->"
        garbled_report = GarbleReport(is_garbled=True, fired_prongs=frozenset({"test"}))
        monkeypatch.setattr(helpers, "detect_garble", lambda text, **kw: garbled_report)
        with patch.dict(
            sys.modules, {"pypdfium2": _fake_pdfium_module(["þÿ\x02\x01 ¤¤¤ \x03\x04 ÿþ" * 20])}
        ):
            assert _document_level_text_fallback(md, "/fake.pdf") == md

        def _raise(path):
            raise RuntimeError("pdfium open failed")

        broken = types.ModuleType("pypdfium2")
        broken.PdfDocument = _raise
        with patch.dict(sys.modules, {"pypdfium2": broken}):
            assert _document_level_text_fallback(md, "/fake.pdf") == md


# ═════════════════════════════════════════════════════════════════════════
# classify_verdict: content-quality guard, leaf concentration
# ═════════════════════════════════════════════════════════════════════════


def test_cat_b_promoted_content_quality_guard():
    """Design Property 5 (D4): promotion to PASS is blocked if
    ``len(flat_text.strip()) < MIN_FLAT_PROMOTION_CHARS`` OR the ratio of
    image-placeholder blocks to total blocks exceeds 0.5, regardless of
    node_count, max_leaf_ratio or garble status.  Doc 21 regression case: 15
    ``<!-- image -->`` blocks, ~210 total chars -- passes node_count /
    leaf-ratio / garble pre-D4 but must no longer be promoted via
    ``cat_b_promoted``.  Zone-1: without gate evaluation
    (validate_result=None) the early structural-OK return may fire with PASS,
    so the key property is that ``cat_b_promoted`` is never the reason.

    Note: ``_flatten_tree_text`` concatenates node text with no separator, so
    per-block text carries a trailing "\\n" here (as real extracted markdown
    blocks do) to make each block land on its own line for the
    placeholder-ratio line-scan in ``classify_verdict``.
    """
    placeholder = [{"title": "", "text": _IMAGE_MARKER + "\n"} for _ in range(15)]
    _verdict, reason = classify_verdict(placeholder, "flat_prose", None)
    assert reason != "cat_b_promoted"

    real = [
        {
            "title": "",
            "text": (
                f"block number {i} has real prose content describing the "
                "document in detail with enough words to be meaningful. " * 3 + "\n"
            ),
        }
        for i in range(15)
    ]
    assert len("".join(b["text"] for b in real).strip()) >= 500
    verdict, reason = classify_verdict(real, "flat_prose", None)
    assert verdict == "PASS"
    assert reason in ("structural_pass", "cat_b_promoted")


def test_pass_max_leaf_ratio_threshold_is_env_tunable_with_widened_default(monkeypatch):
    """Design Property 10 (D10) + D0: the leaf-concentration threshold for the
    main PASS gate reads from PASS_MAX_LEAF_RATIO rather than a hardcoded
    value, and its default is the WIDENED 0.30 (0.25 sits above the OLD 0.20
    default but below the widened one -- the exact regression D0 fixes)."""
    cases = [
        ("0.20", 0.18, ("PASS", "structural_pass")),
        ("0.20", 0.22, ("MARGINAL", "leaf_concentration=0.22")),
        (None, 0.25, ("PASS", "structural_pass")),
        (None, 0.35, ("MARGINAL", "leaf_concentration=0.35")),
    ]
    failures = []
    for env, ratio, want in cases:
        if env is None:
            monkeypatch.delenv("PASS_MAX_LEAF_RATIO", raising=False)
        else:
            monkeypatch.setenv("PASS_MAX_LEAF_RATIO", env)
        reset_pipeline_config()
        got = classify_verdict(_tree_with_ratio(ratio), "hierarchical", None)
        if tuple(got) != want:
            failures.append(
                f"PASS_MAX_LEAF_RATIO={env}, ratio={ratio}: got {tuple(got)}, want {want}"
            )
    assert not failures, "\n".join(failures)


# ═════════════════════════════════════════════════════════════════════════
# Flat-route structure synthesis, VLM fallback, LLM failure classification
# ═════════════════════════════════════════════════════════════════════════


def _synthesize_flat_structure(flat_structure: list, blocks: list) -> list:
    # D5 (RFC-023): mirrors client.py's index() -- always prefer synthetic
    # structure from blocks when blocks exist, regardless of whether
    # flat_structure (the rejected tree) is empty or non-empty.
    if blocks:
        flat_structure = [
            {"title": "", "text": _flat_block_primary_text(b)}
            for b in blocks
            if _flat_block_primary_text(b).strip()
        ]
    return flat_structure


def test_synthetic_structure_preferred_over_rejected_tree():
    """Design Property 6 (D5): for any flat-routed document where ``blocks`` is
    non-empty, the verdict-computation input is the synthetic structure built
    from ``blocks``, regardless of whether the rejected tree structure is
    itself empty or non-empty.  Doc 20 regression case: the tree builder
    produced a non-empty rejected structure (low node_count/depth) but 355 real
    blocks exist -- the rejected structure must never be used.  The pre-D5
    behaviour (structure=[] with blocks) is preserved (no B1/RFC-022
    regression)."""
    rejected = [{"title": "", "text": "sparse rejected tree content"}]
    blocks = [{"text": f"block {i} has real prose content"} for i in range(355)]
    structure = _synthesize_flat_structure(rejected, blocks)
    assert structure != rejected
    assert len(structure) == len(blocks)
    assert all(node["text"] for node in structure)

    small = [{"text": "alpha content"}, {"text": "beta content"}, {"text": "gamma content"}]
    assert len(_synthesize_flat_structure([], small)) == len(small)

    # D6: for a text-only document ``_flat_block_primary_text`` measures the
    # same char count as the pre-fix ``b["text"]`` sum -- the flat-doc char
    # count is unchanged from prior behaviour.
    text_blocks = [
        {"role": "prose", "text": "Clause 1: introductory text."},
        {"role": "prose", "text": "Clause 2: further provisions."},
    ]
    pre_fix_chars = sum(len(b.get("text", "")) for b in text_blocks)
    flat_char_count = sum(len(_flat_block_primary_text(b)) for b in text_blocks)
    assert flat_char_count == pre_fix_chars
    assert flat_char_count == sum(len(b["text"]) for b in text_blocks)


@pytest.mark.asyncio
async def test_vlm_tesseract_fallback_reason_override(monkeypatch):
    """Design Property 8 (D7): on VLM exception, Tesseract OCR runs on the
    rasterized page images; clean OCR text overrides the reason to
    'node_count<3' (flat success path), while garbled output must NOT override
    it -- the document still raises LowQualityTreeError('garbling') per HR5.

    Drives the real ``_attempt_tesseract_raster_recovery``: a non-None return
    is the override signal, None keeps the 'garbling' reason."""
    from pageindex_mcp.client import images as images_mod

    monkeypatch.setattr(images_mod, "detect_ocr_langs", lambda _f: ["eng"])
    monkeypatch.setattr(images_mod, "ensure_tessdata", lambda langs: langs)
    ocr = AsyncMock()
    monkeypatch.setattr(converters_mod, "tesseract_ocr_pdf_pages", ocr)

    ocr.return_value = _CLEAN_TEXT
    assert (
        await images_mod._attempt_tesseract_raster_recovery("/f.pdf", None, "f.pdf") == _CLEAN_TEXT
    )
    ocr.return_value = _GARBLED_TEXT
    assert await images_mod._attempt_tesseract_raster_recovery("/f.pdf", None, "f.pdf") is None


def test_classify_llm_failure_terminal_vs_transient():
    """Design Property 9 (D8): LLMTransientFailure is classified terminal (no
    retry) iff the error detail carries a CMap-corruption or content-policy
    indicator, else transient (retryable)."""
    assert _classify_llm_failure("CMap corruption detected") == "llm_failure_terminal"
    assert _classify_llm_failure("429 rate_limit exceeded, throttled") == "llm_failure_transient"


@pytest.mark.asyncio
async def test_structural_failure_ocr_escalation_for_image_dominant_docs(monkeypatch):
    """Design Property 12 (D11): for any validate_tree failure with reason in
    ('node_count<3', 'depth<2') where the image-line ratio (image lines /
    non-empty lines) exceeds 0.50, the system triggers the same OCR escalation
    path as reason == 'garbling'; the ratio is computed against
    non_empty_lines, not total_lines.  Drives the real
    ``RecoveryMixin._recover_image_dominant_ocr`` with a spy OCR retry."""
    from pageindex_mcp.client import recovery as recovery_mod
    from pageindex_mcp.helpers import TreeDefect, TreeGateResult

    monkeypatch.setattr(
        recovery_mod,
        "pipeline_config",
        dataclasses.replace(pipeline_config, image_dominant_ocr_escalation_enabled=True),
    )
    monkeypatch.setattr(recovery_mod, "settings", SimpleNamespace(flat_doc_routing=True))
    retry = AsyncMock(return_value=True)
    monkeypatch.setattr(recovery_mod.RecoveryMixin, "_execute_ocr_retry", retry)

    def _state(md):
        return ExtractionState(
            result={"structure": []},
            ok=False,
            reason="node_count<3",
            gate_result=TreeGateResult(ok=False, defect=TreeDefect.NODE_COUNT_LOW),
            first_defect=TreeDefect.NODE_COUNT_LOW,
            route=Route.FLAT,
            md_content=md,
            tmp_md_path=None,
            pic_results=[],
            used_converter="docling",
            total_chars=len(md),
            extraction_stages_captured=[],
        )

    mixin = recovery_mod.RecoveryMixin()
    # 3/4 non-empty lines are image markers; blank lines must not dilute it.
    dominant = _state(f"{_MARKER}\n\n\n{_MARKER}\n\n{_MARKER}\n\nsome prose")
    await mixin._recover_image_dominant_ocr(dominant, "/f.pdf", "f.pdf", ".pdf", None)
    assert retry.await_count == 1
    assert dominant.full_page_already_applied is True

    sparse = _state("\n".join(["real paragraph text here"] * 8 + [_MARKER]))
    await mixin._recover_image_dominant_ocr(sparse, "/f.pdf", "f.pdf", ".pdf", None)
    assert retry.await_count == 1, "1/9 image lines must not escalate"
    assert sparse.full_page_already_applied is False


# ═════════════════════════════════════════════════════════════════════════
# pipeline.py: OCR recovery cascade — marker cleanup, decide_ocr_mode removal
# ═════════════════════════════════════════════════════════════════════════


def test_image_markers_stripped_only_when_pic_results_empty(monkeypatch):
    """Regression: when ``_recover_picture_results`` returns [] (OCR skip),
    ``_fallback_and_recover_pictures`` strips residual ``<!-- image -->``
    markers from the markdown it returns; when pic_results ARE populated the
    markers are preserved as splice targets for bind_markers."""
    from pageindex_mcp.converters import pipeline as pipeline_mod

    monkeypatch.setattr(converters.pictures, "_recover_picture_results", lambda *a, **kw: [])
    md_out, pic_results, _records = pipeline_mod._fallback_and_recover_pictures(
        "# Heading\n\n<!-- image -->\n\nBody text <!-- image --> end",
        document=None,
        pdf_path="/fake.pdf",
        filename="fake.pdf",
        expected_script=None,
        landscape_fallback_pages=[],
        heading_pages={},
        force_full_page_ocr_applied=True,  # force OCR skip
    )
    assert "<!-- image -->" not in md_out, (
        "Residual <!-- image --> markers must be stripped when pic_results is empty"
    )
    assert pic_results == []

    fake_pr = {"ocr_text": "chart", "page": 1, "bbox": {}, "png_bytes": b"png"}
    monkeypatch.setattr(pipeline_mod, "_recover_picture_results", lambda *a, **kw: [fake_pr])
    _md_out, pic_results, _records = pipeline_mod._fallback_and_recover_pictures(
        "# H\n\n<!-- image -->\n\nBody",
        document=None,
        pdf_path="/fake.pdf",
        filename="fake.pdf",
        expected_script=None,
        landscape_fallback_pages=[],
        heading_pages={},
        force_full_page_ocr_applied=False,
    )
    assert len(pic_results) == 1


def test_decide_ocr_mode_removed_and_strategy_is_canonical():
    """Wiring: the ``decide_ocr_mode`` wrapper is removed from picture_plane and
    is not re-exported from converters; ``decide_ocr_strategy`` is the
    canonical replacement and ``converters/pictures.py`` no longer defines the
    wrapper."""
    import inspect

    from pageindex_mcp import converters as converters_pkg, picture_plane
    from pageindex_mcp.converters import pictures
    from pageindex_mcp.picture_plane import decide_ocr_strategy

    assert not hasattr(picture_plane, "decide_ocr_mode")
    assert not hasattr(converters_pkg, "decide_ocr_mode")
    assert inspect.isfunction(decide_ocr_strategy)
    assert "def decide_ocr_mode" not in inspect.getsource(pictures)


def test_verdict_ledger_helpers_removed_from_storage():
    """RFC-037 D3: ``persist_verdict_ledger`` / ``read_verdict_ledger`` must no
    longer be importable, and ``read_verdict_ledger`` must be gone from
    ``storage.__all__``."""
    import pageindex_mcp.storage as storage_mod

    assert not hasattr(storage_mod, "persist_verdict_ledger")
    assert not hasattr(storage_mod, "read_verdict_ledger")
    assert "read_verdict_ledger" not in storage_mod.__all__


# ═════════════════════════════════════════════════════════════════════════
# helpers pipeline: apply_promotions / evaluate_gates / verdict authority
# ═════════════════════════════════════════════════════════════════════════


def _th() -> VerdictThresholds:
    return VerdictThresholds.from_config(pipeline_config)


def _well_formed() -> list:
    return [
        {
            "node_id": "1",
            "title": "Root",
            "text": "",
            "nodes": [
                {"node_id": "2", "title": "Ch1", "text": "a" * 100, "nodes": []},
                {"node_id": "3", "title": "Ch2", "text": "b" * 100, "nodes": []},
                {"node_id": "4", "title": "Ch3", "text": "c" * 100, "nodes": []},
            ],
        }
    ]


def _varied_text(i: int) -> str:
    paragraphs = [
        "The insurance contract shall be governed by the applicable laws and regulations.",
        "Premium payments are due on the first day of each calendar month without exception.",
        "Coverage extends to all listed beneficiaries as specified in the policy document.",
    ]
    return paragraphs[i % len(paragraphs)]


def _outcome_for(
    structure: list | None = None,
    defect: TreeDefect = TreeDefect.OK,
    all_defects: frozenset | None = None,
) -> GateOutcome:
    if structure is None:
        structure = _well_formed()
    th = _th()
    sig = TreeSignals.from_tree(structure, garble_threshold=th.garble_threshold)
    return GateOutcome(
        defect=defect,
        validate_reason=None,
        signals=sig,
        all_defects=all_defects if all_defects is not None else frozenset(),
        hard_fail_verdict=None,
    )


def _make_gate_result(
    defect: TreeDefect,
    structure: list | None = None,
    all_defects: frozenset | None = None,
) -> TreeGateResult:
    if structure is None:
        structure = _well_formed()
    sig = TreeSignals.from_tree(structure, garble_threshold=_th().garble_threshold)
    if all_defects is None:
        all_defects = frozenset({defect}) if defect != TreeDefect.OK else frozenset()
    return TreeGateResult(
        ok=(defect == TreeDefect.OK),
        defect=defect,
        detail=defect.value,
        signals=sig,
        all_defects=all_defects,
    )


def test_apply_promotions_passes_well_formed_and_high_enrichment_images():
    """A well-formed multi-chapter tree PASSes on its own structure; an
    image_standalone document with a 0.95 enrichment ratio is promoted to
    PASS."""
    structure = [
        {
            "node_id": "1",
            "title": "Root",
            "text": "",
            "nodes": [
                {"node_id": str(i), "title": f"Chapter {i}", "text": _varied_text(i), "nodes": []}
                for i in range(2, 12)
            ],
        }
    ]
    vr = apply_promotions(
        _outcome_for(structure=structure),
        "",
        image_enrichment_ratio=None,
        inspector_class=None,
        th=_th(),
        expected_script=None,
    )
    assert vr.verdict == "PASS"
    verdict, _reason = vr  # VerdictResult stays tuple-unpackable
    assert verdict == "PASS"

    vr = apply_promotions(
        _outcome_for(),
        "image_standalone",
        image_enrichment_ratio=0.95,
        inspector_class=None,
        th=_th(),
        expected_script=None,
    )
    assert vr.verdict == "PASS"


def test_evaluate_gates_hard_fail_verdict_and_cofiring_tiebreak():
    """A non-hard-fail defect (and TreeDefect.OK) yields no hard_fail_verdict;
    when two hard-fail defects co-fire the reason is the one with the worst
    ``_GATE_PRIORITY``."""
    outcome = evaluate_gates(
        _well_formed(), _make_gate_result(TreeDefect.NODE_COUNT_LOW), None, _th()
    )
    assert outcome.hard_fail_verdict is None

    outcome = evaluate_gates(_well_formed(), _make_gate_result(TreeDefect.OK), None, _th())
    assert outcome.hard_fail_verdict is None
    assert outcome.defect == TreeDefect.OK

    hf_list = sorted(HARD_FAIL_DEFECTS, key=lambda d: _GATE_PRIORITY.get(d, 999))
    if len(hf_list) < 2:
        pytest.skip("Need at least 2 hard-fail defects")
    worst, second = hf_list[0], hf_list[1]
    gr = _make_gate_result(TreeDefect.OK, all_defects=frozenset({worst, second}))
    outcome = evaluate_gates(_well_formed(), gr, None, _th())
    assert outcome.hard_fail_verdict is not None
    assert outcome.hard_fail_verdict.reason == worst.value


@pytest.mark.asyncio
async def test_upsert_verdict_returns_winning_row():
    from pageindex_mcp.registry import upsert_verdict

    winning = {
        "doc_id": "abc",
        "verdict": "PASS",
        "pipeline_version": 4,
        "permanent_marginal": False,
        "verdict_computed_at": "2026-08-18T12:00:00Z",
    }
    mock_pool = AsyncMock()
    mock_pool.fetchrow = AsyncMock(return_value=winning)
    with patch("pageindex_mcp.registry.schema.get_pool", return_value=mock_pool):
        result = await upsert_verdict(
            "abc", {"verdict": "PASS", "verdict_computed_at": "2026-08-18T12:00:00Z"}
        )
    assert result["verdict"] == "PASS"


def test_run_stages_records_order_and_isolates_stage_failures():
    """``_run_stages`` records one provenance entry per stage in order, and a
    stage that raises does not skip the next one -- its error is recorded and
    the pipeline continues from the last good markdown."""
    from pageindex_mcp.converters import _run_stages

    _md, records = _run_stages("x", [("alpha", lambda m: m), ("beta", lambda m: m + "!")])
    assert list(records.keys()) == ["alpha", "beta"]

    def fail(md):
        raise RuntimeError("boom")

    md, records = _run_stages("start", [("fail", fail), ("ok", lambda m: m + " ok")])
    assert md == "start ok"
    assert records["fail"]["error"] is not None
    assert records["ok"]["error"] is None


def test_prepare_tree_composes_split_then_segment():
    """A small structure passes through ``prepare_tree`` unchanged; a large one
    is exactly ``_segment_table_nodes(split_oversized_leaf_nodes(...))``."""
    small = [
        {
            "title": "S1",
            "text": "Short.",
            "level": 1,
            "nodes": [{"title": "Sub", "text": "Details.", "level": 2}],
        },
    ]
    assert prepare_tree(copy.deepcopy(small)) == small

    big_text = "\n\n".join(f"Article ({i})\n\n" + "Body. " * 4000 for i in range(1, 5))
    structure = [{"title": "Doc", "text": big_text, "level": 1}]
    result = prepare_tree(copy.deepcopy(structure))
    manual = _segment_table_nodes(split_oversized_leaf_nodes(copy.deepcopy(structure)))
    assert result == manual


def test_tag_landscape_pages_for_fallback(tmp_path):
    fitz = pytest.importorskip("fitz")
    from pageindex_mcp.converters import _tag_landscape_pages_for_fallback

    doc = fitz.open()
    doc.new_page(width=600, height=800)
    path = str(tmp_path / "portrait.pdf")
    doc.save(path)
    doc.close()
    pages = _tag_landscape_pages_for_fallback(path)
    assert pages[0]["is_landscape"] is False


def test_build_candidate_matches_the_mirrored_stage_sequence():
    """``_build_candidate`` is exactly the Arabic -> German -> English heading
    injections followed by ``_pre_inference_normalize``."""
    from pageindex_mcp.converters import (
        _build_candidate,
        _pre_inference_normalize,
    )

    def _mirrored(md):
        md = _inject_arabic_structural_headings(md)
        md = _inject_german_clause_headings(md)
        md = _inject_english_article_headings(md)
        return _pre_inference_normalize(md)

    assert _build_candidate("") == _mirrored("")


# ═════════════════════════════════════════════════════════════════════════
# formats.py: tesseract_ocr_pdf_pages dual rasterization backend (D4)
# ═════════════════════════════════════════════════════════════════════════

_PDFIUM_PNG = f"data:image/png;base64,{base64.b64encode(b'PDFIUM_PNG_FAKE').decode()}"
_FITZ_PNG = f"data:image/png;base64,{base64.b64encode(b'FITZ_PNG_FAKE').decode()}"


async def test_tesseract_ocr_pdf_pages_fitz_fallback_gating(monkeypatch):
    """When pypdfium2 rasterization succeeds the fitz fallback is never called;
    when the fallback is disabled a pypdfium2 failure propagates untouched
    (fitz still not called)."""
    fitz_called = False

    def _fitz_fallback(pdf_path, dpi=200):
        nonlocal fitz_called
        fitz_called = True
        return [_FITZ_PNG]

    monkeypatch.setattr(
        converters.formats, "rasterize_pdf_pages", lambda pdf_path, dpi=200: [_PDFIUM_PNG]
    )
    monkeypatch.setattr(converters.formats, "rasterize_pdf_pages_fitz", _fitz_fallback)
    monkeypatch.setattr(
        converters.pictures, "_tesseract_ocr_image", lambda path, langs: "pdfium text"
    )

    assert await converters.tesseract_ocr_pdf_pages("/fake.pdf", ["eng"]) == "pdfium text"
    assert fitz_called is False

    monkeypatch.setattr(converters.formats, "_D7_FITZ_FALLBACK_ENABLED", False)

    def _pdfium_boom(pdf_path, dpi=200):
        raise RuntimeError("CMap corruption: pypdfium2 render failed")

    monkeypatch.setattr(converters.formats, "rasterize_pdf_pages", _pdfium_boom)
    with pytest.raises(RuntimeError, match="CMap corruption"):
        await converters.tesseract_ocr_pdf_pages("/fake.pdf", ["eng"])
    assert fitz_called is False


# ═════════════════════════════════════════════════════════════════════════
# split_oversized_leaf_nodes / _has_heading_markers / _ordinal_value (D3)
# ═════════════════════════════════════════════════════════════════════════


class TestSplitOversizedLeafNodes:
    def test_ordinal_markers_detected_and_split_fires(self):
        """'Clause N' is detected by ``_has_heading_markers`` (split-eligible
        even under max_chars) and an oversized leaf carrying those markers
        splits into 3 children; the pre-existing Article/Section/مادة patterns
        still match and still split (no regression from the new marker types);
        roman numerals ('Part IV/V/VI') and letters ('Annex A/B/C') resolve to
        the correct ``_ordinal_value`` int tuples."""
        assert _has_heading_markers("Clause 1 says X. Clause 2 says Y. Clause 3 says Z.") is True

        failures = []
        for marker in ("Article 9", "Section 4", "المادة ٥"):
            if _OVERSIZED_ORDINAL_RE.search(marker) is None:
                failures.append(f"_OVERSIZED_ORDINAL_RE does not match {marker!r}")
        ordinal_cases = [
            ("Part IV", (4,)),
            ("Part V", (5,)),
            ("Part VI", (6,)),
            ("Annex A", (1,)),
            ("Annex B", (2,)),
            ("Annex C", (3,)),
        ]
        for text, want in ordinal_cases:
            got = _ordinal_value(_OVERSIZED_ORDINAL_RE.search(text))
            if got != want:
                failures.append(f"_ordinal_value({text!r}) == {got}, want {want}")

        for prefix in ("Clause", "Article"):
            text = "\n".join(f"{prefix} {i} {_text_of_length(3000)}" for i in (1, 2, 3))
            tree = [{"node_id": "n1", "title": "root", "text": text, "nodes": []}]
            split_oversized_leaf_nodes(tree, max_chars=50000, min_segments=3)
            if len(tree[0]["nodes"]) != 3:
                failures.append(
                    f"{prefix}: split produced {len(tree[0]['nodes'])} children, want 3"
                )
            elif not tree[0]["nodes"][0]["text"].startswith(f"{prefix} 1"):
                failures.append(f"{prefix}: first child does not start at '{prefix} 1'")
        assert not failures, "\n".join(failures)

    def test_part_prose_false_positive_regression_guard(self):
        """'Part 2 of the agreement' repeated (non-sequential, same ordinal
        each time) is English prose making a cross-reference, not a heading
        sequence -> must NOT produce a spurious split."""
        text = (
            f"As mentioned in Part 2 of the agreement, {_text_of_length(2500)}\n\n"
            f"Part 2 of the agreement also states {_text_of_length(2500)}\n\n"
            f"Referring again to Part 2 of the agreement, {_text_of_length(2500)}"
        )
        tree = [{"node_id": "n1", "title": "root", "text": text, "nodes": []}]
        split_oversized_leaf_nodes(
            tree, max_chars=50000, min_segments=3, _tree_ratio=0.1, _tree_total=len(text) * 10
        )
        assert tree[0]["nodes"] == []


# ═════════════════════════════════════════════════════════════════════════
# converters_cli: `python -m pageindex_mcp.converters_cli <input_pdf_path>`
# ═════════════════════════════════════════════════════════════════════════
#
# Contract under test:
#     stdout: the RFC-028 D0 startup handshake line followed by exactly one
#             result JSON line at exit
#       success: {"ok": true, "doc_id": ..., "peak_rss_kib": int, "duration_ms": int}
#       failure: {"ok": false, "error": "<ExceptionClassName>", "message": ...}
#     exit code: 0 on success, 1 on handled exception.

# Minimal valid PDF bytes (no text, just structure — enough for file-exists checks).
_MINIMAL_PDF = (
    b"%PDF-1.4\n"
    b"1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
    b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
    b"3 0 obj<</Type/Page/MediaBox[0 0 612 792]/Parent 2 0 R>>endobj\n"
    b"xref\n0 4\n"
    b"0000000000 65535 f \n"
    b"0000000009 00000 n \n"
    b"0000000058 00000 n \n"
    b"0000000115 00000 n \n"
    b"trailer<</Size 4/Root 1 0 R>>\n"
    b"startxref\n190\n%%EOF"
)


@pytest.fixture()
def tmp_pdf(tmp_path: Path) -> Path:
    """Write a minimal PDF fixture to a temp file and return its path."""
    p = tmp_path / "fixture.pdf"
    p.write_bytes(_MINIMAL_PDF)
    return p


def _run_cli(*args, env_extra=None, timeout=180):
    """Run the CLI as a subprocess and return the CompletedProcess."""
    env = os.environ.copy()
    if env_extra:
        env.update(env_extra)
    return subprocess.run(
        [sys.executable, "-m", "pageindex_mcp.converters_cli", *args],
        capture_output=True,
        timeout=timeout,
        env=env,
    )


def _last_stdout_json(proc) -> dict:
    """Return the last non-empty stdout line parsed as JSON."""
    lines = [ln for ln in proc.stdout.decode().splitlines() if ln.strip()]
    assert lines, f"No stdout output. stderr={proc.stderr.decode()!r}"
    return json.loads(lines[-1])


def _neutralize_llm_gate(monkeypatch):
    """The CLI validates/configures the LLM provider (LLM-01) before indexing.
    These tests exercise index() plumbing, not provider config, so the gate is
    neutralized to stay independent of ambient OPENAI_API_KEY in the runner."""
    monkeypatch.setattr("pageindex_mcp.client.llm.validate_llm_config", lambda: None)
    monkeypatch.setattr("pageindex_mcp.client.llm.configure_litellm", lambda: None)
    monkeypatch.setattr("pageindex_mcp.client.validate_llm_config", lambda: None)
    monkeypatch.setattr("pageindex_mcp.client.configure_litellm", lambda: None)


def test_cli_missing_input_file_exits_1_with_json_error(tmp_path: Path):
    """Missing input → exit code 1, JSON {ok: false, error: FileNotFoundError}."""
    proc = _run_cli(str(tmp_path / "does_not_exist.pdf"), timeout=30)
    assert proc.returncode == 1, f"Expected exit 1. stderr={proc.stderr.decode()!r}"
    payload = _last_stdout_json(proc)
    assert payload["ok"] is False
    assert payload["error"] == "FileNotFoundError"
    assert "message" in payload


async def test_cli_runtime_error_from_index_exits_1(tmp_pdf: Path, monkeypatch):
    """RuntimeError raised by client.index → exit 1 with
    JSON {ok: false, error: RuntimeError} carrying the message."""
    import pageindex_mcp.converters_cli as cli_module
    from pageindex_mcp.converters_cli import main

    monkeypatch.setattr("sys.argv", ["converters_cli", str(tmp_pdf)])
    fake_stdout = io.StringIO()
    monkeypatch.setattr(cli_module, "_stdout", fake_stdout)
    _neutralize_llm_gate(monkeypatch)

    with patch(
        "pageindex_mcp.client.CustomPageIndexClient.index",
        new_callable=AsyncMock,
        side_effect=RuntimeError("empty pdf"),
    ):
        exit_code = await main()

    assert exit_code == 1
    output = fake_stdout.getvalue().strip()
    assert output, "Expected one stdout line"
    payload = json.loads(output.splitlines()[-1])
    assert payload["ok"] is False
    assert payload["error"] == "RuntimeError"
    assert "empty pdf" in payload["message"]


async def test_cli_success_json_shape_and_handshake(tmp_pdf: Path, monkeypatch):
    """RFC-028 D0: exactly 2 stdout lines — the startup handshake
    ({"handshake": true, ...}) followed by the result JSON, whose keys are
    exactly {ok, doc_id, peak_rss_kib, duration_ms} with the right types."""
    import pageindex_mcp.converters_cli as cli_module
    from pageindex_mcp.converters_cli import main

    monkeypatch.setattr("sys.argv", ["converters_cli", str(tmp_pdf)])
    fake_stdout = io.StringIO()
    monkeypatch.setattr(cli_module, "_stdout", fake_stdout)
    _neutralize_llm_gate(monkeypatch)

    with patch(
        "pageindex_mcp.client.CustomPageIndexClient.index",
        new_callable=AsyncMock,
        return_value="deadbeef",
    ):
        exit_code = await main()

    assert exit_code == 0
    lines = [ln for ln in fake_stdout.getvalue().splitlines() if ln.strip()]
    assert len(lines) == 2, f"Expected exactly 2 stdout lines, got: {lines}"
    handshake = json.loads(lines[0])
    assert handshake["handshake"] is True
    assert handshake["chunk_count"] == 1
    assert handshake["is_docling_route"] is True

    payload = json.loads(lines[1])
    assert set(payload.keys()) == {"ok", "doc_id", "peak_rss_kib", "duration_ms"}
    assert payload["ok"] is True
    assert isinstance(payload["doc_id"], str) and len(payload["doc_id"]) > 0
    assert isinstance(payload["peak_rss_kib"], int) and payload["peak_rss_kib"] >= 0
    assert isinstance(payload["duration_ms"], int) and payload["duration_ms"] >= 0


def test_cli_stdout_not_polluted_by_logs_or_stray_prints(tmp_pdf: Path, tmp_path: Path):
    """Real-subprocess contract: stdout is exactly the handshake line plus the
    result JSON even when client.index prints to stdout and logs a warning."""
    shim = tmp_path / "noisy_shim.py"
    shim.write_text(
        "import sys, logging, asyncio\n"
        "from unittest.mock import patch\n"
        "sys.argv = ['converters_cli', sys.argv[1]]\n"
        "\n"
        "async def noisy_index(self_or_path, *a, **kw):\n"
        "    print('noisy stdout log')  # stray print — should go to stderr or be suppressed\n"
        "    logging.getLogger().warning('noisy warning')\n"
        "    return 'cafe5678'\n"
        "\n"
        "with patch('pageindex_mcp.client.CustomPageIndexClient.index', noisy_index):\n"
        "    from pageindex_mcp.converters_cli import main\n"
        "    sys.exit(asyncio.run(main()))\n"
    )
    result = subprocess.run(
        [sys.executable, str(shim), str(tmp_pdf)], capture_output=True, timeout=60
    )
    stdout_lines = [ln for ln in result.stdout.decode().splitlines() if ln.strip()]
    assert len(stdout_lines) == 2, (
        f"Expected exactly 2 stdout lines, got {len(stdout_lines)}: {stdout_lines!r}. "
        f"stderr={result.stderr.decode()!r}"
    )
    payload = json.loads(stdout_lines[-1])
    assert payload["ok"] is True
    assert payload["doc_id"] == "cafe5678"


async def test_cli_registry_fields_surfaced_iff_client_stashed_them(tmp_pdf: Path, monkeypatch):
    """Zone-7 dual-write consistency: when ``client.last_registry_fields`` is
    set after index(), the stdout JSON carries a ``registry_fields`` dict with
    every ``_REGISTRY_FIELDS`` key (minus doc_id) plus node_count; when it is
    None the key must be absent (backward compat)."""
    import pageindex_mcp.converters_cli as cli_module
    from pageindex_mcp.converters_cli import main
    from pageindex_mcp.storage.verdict import _REGISTRY_FIELDS

    expected_keys = {k for k in _REGISTRY_FIELDS if k != "doc_id"} | {"node_count"}
    stashed = {k: "" for k in expected_keys}
    stashed.update(
        {
            "doc_name": "fixture.pdf",
            "source_url": "http://x",
            "processed_at": "2026-08-26T00:00:00Z",
            "sha256": "abc123",
            "content_class": "flat_prose",
            "node_count": 0,
        }
    )

    async def _fake_index(self, *a, **kw):
        self.last_registry_fields = stashed
        return "reg-fields-1"

    monkeypatch.setattr("sys.argv", ["converters_cli", str(tmp_pdf)])
    fake_stdout = io.StringIO()
    monkeypatch.setattr(cli_module, "_stdout", fake_stdout)
    _neutralize_llm_gate(monkeypatch)
    with patch("pageindex_mcp.client.CustomPageIndexClient.index", _fake_index):
        exit_code = await main()

    assert exit_code == 0
    payload = json.loads([ln for ln in fake_stdout.getvalue().splitlines() if ln.strip()][-1])
    assert payload["ok"] is True
    rf = payload.get("registry_fields")
    assert rf is not None, "registry_fields must be surfaced when the client stashed them"
    missing = expected_keys - set(rf)
    assert not missing, f"Missing registry fields: {sorted(missing)}"

    fake_stdout = io.StringIO()
    monkeypatch.setattr(cli_module, "_stdout", fake_stdout)
    with patch(
        "pageindex_mcp.client.CustomPageIndexClient.index",
        new_callable=AsyncMock,
        return_value="no-reg-1",
    ):
        exit_code = await main()
    assert exit_code == 0
    payload = json.loads([ln for ln in fake_stdout.getvalue().splitlines() if ln.strip()][-1])
    assert payload["ok"] is True
    assert "registry_fields" not in payload


@pytest.mark.integration
def test_cli_happy_path_real_docling_integration(tmp_pdf: Path):
    """Real Docling conversion — only runs when DOCLING_INTEGRATION=1."""
    if not os.environ.get("DOCLING_INTEGRATION"):
        pytest.skip("Set DOCLING_INTEGRATION=1 to run real Docling integration test")
    proc = _run_cli(str(tmp_pdf), timeout=300)
    assert proc.returncode == 0, f"CLI failed. stderr={proc.stderr.decode()!r}"
    payload = _last_stdout_json(proc)
    assert payload["ok"] is True
    assert isinstance(payload["doc_id"], str) and len(payload["doc_id"]) > 0
    assert isinstance(payload["peak_rss_kib"], int) and payload["peak_rss_kib"] >= 0
    assert isinstance(payload["duration_ms"], int) and payload["duration_ms"] >= 0


# ═════════════════════════════════════════════════════════════════════════
# RFC-041 D10c: pre-NFKC ScriptContext threading
# ═════════════════════════════════════════════════════════════════════════
#
# NFKC decomposition destroys Arabic Presentation Forms codepoints, so a
# ScriptContext built from post-NFKC text always reports
# had_presentation_forms=False.  Every garble call site must therefore accept a
# caller-supplied ScriptContext and thread it through INSTEAD of re-inferring
# from the normalized text -- otherwise Arabic documents lose their PF signal.

# Arabic Presentation Forms-B text (U+FE70-U+FEFF range)
_ARABIC_PF_TEXT = "ﺍﺎﺻﺼﺵﺶ"
_ARABIC_PF_NFKC = unicodedata.normalize("NFKC", _ARABIC_PF_TEXT)

# Latin text — no presentation forms in either form
_LATIN_TEXT = "The quick brown fox jumps over the lazy dog"


def test_infer_presentation_forms_is_destroyed_by_nfkc():
    """``_infer_presentation_forms`` sees Arabic Presentation Forms pre-NFKC and
    loses them post-NFKC; Latin text answers False either way (parity)."""
    from pageindex_mcp.helpers.garble import _infer_presentation_forms

    assert _infer_presentation_forms(_ARABIC_PF_TEXT) is True
    assert _infer_presentation_forms(_ARABIC_PF_NFKC) is False

    pre = _infer_presentation_forms(_LATIN_TEXT)
    post = _infer_presentation_forms(unicodedata.normalize("NFKC", _LATIN_TEXT))
    assert pre == post
    assert pre is False


def test_script_context_pf_is_threaded_to_every_garble_call_site(monkeypatch):
    """A caller-supplied ScriptContext reaches ``detect_garble`` with its own
    ``had_presentation_forms`` at every threaded call site -- ``_garble_ratio``,
    ``TreeSignals.from_tree``, ``validate_tree`` and ``compute_verdict`` --
    instead of the value re-inferred from the (post-NFKC) text, which would
    always be False.  This is the regression D10c fixes."""
    from pageindex_mcp.helpers import garble as garble_mod
    from pageindex_mcp.helpers.garble import GarbleReport, _garble_ratio
    from pageindex_mcp.helpers.tree_validation import TreeSignals as _TreeSignals, validate_tree
    from pageindex_mcp.helpers.verdict import compute_verdict
    from pageindex_mcp.script import ScriptContext

    seen: list[bool | None] = []

    def _spy(text, *, script_context=None, **kw):
        seen.append(script_context.had_presentation_forms if script_context else None)
        return GarbleReport(is_garbled=False, fired_prongs=frozenset())

    monkeypatch.setattr(garble_mod, "detect_garble", _spy)

    flat_structure = [{"heading": "test", "content": _ARABIC_PF_NFKC, "children": []}]
    nested_structure = [
        {
            "heading": "root",
            "content": "",
            "children": [{"heading": "child", "content": _ARABIC_PF_NFKC, "children": []}],
        }
    ]

    failures = []
    for pf in (True, False):
        ctx = ScriptContext(
            dominant_script="Arab",
            had_presentation_forms=pf,
            source=f"test_pf_{pf}",
        )
        sites = {
            "_garble_ratio": lambda: _garble_ratio(
                _ARABIC_PF_NFKC, expected_script="Arab", script_context=ctx
            ),
            "TreeSignals.from_tree": lambda: _TreeSignals.from_tree(
                flat_structure, expected_script=ctx
            ),
            "validate_tree": lambda: validate_tree(flat_structure, expected_script=ctx),
            "compute_verdict": lambda: compute_verdict(nested_structure, "", expected_script=ctx),
        }
        for name, call in sites.items():
            seen.clear()
            call()
            if not seen:
                failures.append(f"{name} (pf={pf}): detect_garble was never reached")
            elif any(v is not pf for v in seen):
                failures.append(f"{name} (pf={pf}): detect_garble saw {seen}, want all {pf}")
    assert not failures, "\n".join(failures)

    # Latin documents are unaffected by the threading.
    latin_ctx = ScriptContext(
        dominant_script="Latn", had_presentation_forms=False, source="test_latin"
    )
    seen.clear()
    validate_tree(
        [{"heading": "test", "content": _LATIN_TEXT, "children": []}], expected_script=latin_ctx
    )
    assert seen and all(v is False for v in seen)


def test_script_context_from_document_pf_detection_and_enrichment():
    """``ScriptContext.from_document`` on post-NFKC text returns
    had_presentation_forms=False and on pre-NFKC text returns True; the indexer
    enriches the post-NFKC context via ``dataclasses.replace`` without
    disturbing the other fields."""
    from pageindex_mcp.script import ScriptContext

    assert (
        ScriptContext.from_document("arabic.pdf", raw_text=_ARABIC_PF_TEXT).had_presentation_forms
        is True
    )

    ctx = ScriptContext.from_document("arabic.pdf", raw_text=_ARABIC_PF_NFKC)
    assert ctx.had_presentation_forms is False
    enriched = dataclasses.replace(ctx, had_presentation_forms=True)
    assert enriched.had_presentation_forms is True
    assert enriched.dominant_script == ctx.dominant_script
    assert enriched.source == ctx.source


def test_plan_docling_sizes_from_cgroup_limits(monkeypatch):
    """``plan_docling`` derives processes, threads and chunk size from the
    container's CPU and memory, and ``available_*`` read the cgroup v2 limits
    (quota rounded down, "max" meaning unlimited)."""
    from pageindex_mcp.converters import docling_resources as dr

    gib = 1024**3
    rows = [
        # pages, cpus, memory -> workers, threads, pages_per_chunk
        (292, 4, 6783 * 1024**2, 4, 1, 10),  # cx33 pod sized to the node
        (292, 4, 6 * gib, 3, 1, 20),  # memory, not CPUs, caps the processes
        (292, 2, 3 * gib, 1, 2, 36),  # one process: chunked to fit memory
        (292, 8, 15 * gib, 8, 1, 19),  # an 8-vCPU node when cx33 is out of stock
        (292, 14, 64 * gib, 14, 1, 11),  # Mac mini M4 Pro, native (no cgroup)
        (58, 4, 6783 * 1024**2, 1, 4, 58),  # below PARALLEL_MIN_PAGES: one pass
        (5, 4, 6 * gib, 1, 4, 5),
    ]
    bad = []
    for pages, cpus, mem, *want in rows:
        p = dr.plan_docling(pages, cpus=cpus, memory_bytes=mem)
        if [p.workers, p.threads_per_worker, p.pages_per_chunk] != want:
            bad.append((pages, cpus, mem, p))
    assert not bad, bad

    files = {
        "/proc/meminfo": "MemTotal: 8000000 kB",
        "/sys/fs/cgroup/cpu.max": "250000 100000",
        "/sys/fs/cgroup/memory.max": "3221225472",
    }
    monkeypatch.setattr(dr, "_read", files.get)
    monkeypatch.setattr(dr.os, "sched_getaffinity", lambda _pid: set(range(8)), raising=False)
    assert (dr.available_cpus(), dr.available_memory_bytes()) == (2, 3 * gib)
    files.update({"/sys/fs/cgroup/cpu.max": "max 100000", "/sys/fs/cgroup/memory.max": "max"})
    assert (dr.available_cpus(), dr.available_memory_bytes()) == (8, 8000000 * 1024)
    # macOS: no /proc and no cgroup files, so the host's physical memory.
    files.clear()
    pages = {"SC_PHYS_PAGES": 4 * 1024**2, "SC_PAGE_SIZE": 16384}
    monkeypatch.setattr(dr.os, "sysconf", pages.__getitem__)
    assert dr.available_memory_bytes() == 64 * gib


def test_chunked_docling_runs_chunks_in_parallel_in_page_order(tmp_path, monkeypatch):
    """With ``workers`` > 1 the chunks convert concurrently, each child gets
    ``num_threads``, and the markdown and picture pages come back in page
    order whatever order the chunks finish in."""
    import threading
    import time

    fitz = pytest.importorskip("fitz")
    from pageindex_mcp.converters import docling_conv

    doc = fitz.open()
    for _ in range(25):
        doc.new_page()
    path = str(tmp_path / "big.pdf")
    doc.save(path)
    doc.close()

    lock = threading.Lock()
    running = {"now": 0, "peak": 0}
    seen_threads = []

    def fake_chunk(chunk_path, *, num_threads, **_kw):
        with fitz.open(chunk_path) as chunk:
            n = chunk.page_count
        with lock:
            running["now"] += 1
            running["peak"] = max(running["peak"], running["now"])
            seen_threads.append(num_threads)
        time.sleep(0.2 if n == 10 else 0.0)  # full chunks finish after the short last one
        with lock:
            running["now"] -= 1
        return f"<{n}>", [{"page": 1}], {}

    monkeypatch.setattr(docling_conv, "_run_docling_chunk_with_timeout", fake_chunk)
    md, pics, _ = docling_conv._pdf_to_markdown_docling_chunked(
        path, page_count=25, max_pages=10, workers=3, num_threads=2
    )
    assert md == "<10>\n\n<10>\n\n<5>"
    assert [p["page"] for p in pics] == [1, 11, 21]
    assert running["peak"] == 3 and seen_threads == [2, 2, 2]

    # The fake above skips the child; the child's own call must fit the real
    # pipeline signature (it once passed a kwarg that only failed on a Mac run).
    import inspect
    import queue

    from pageindex_mcp.converters import pipeline

    real = inspect.signature(pipeline.pdf_to_markdown_docling)
    tables = []

    def fake_pipeline(*a, **kw):
        real.bind(*a, **kw)
        tables.append(kw.get("pages_with_tables"))
        return "md", [], {}

    monkeypatch.setattr(pipeline, "pdf_to_markdown_docling", fake_pipeline)
    for on in (True, False):
        q = queue.Queue()
        docling_conv._docling_chunk_worker(q, path, False, None, do_table_structure=on)
        assert q.get_nowait() == ("ok", ("md", [], {}))
    assert tables == [None, set()]
